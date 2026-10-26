#!/usr/bin/env python3
"""
Analyze QUAFFLE experiment results across datasets, methods, partitions, and seeds.

Usage:
    python analyze_results.py [--include-backend pennylane [torch ...] | all]

Quantum runs are filtered by the circuit simulator that produced them
(q_backend). By default only q_backend='pennylane' runs enter the tables and
plots; --include-backend adds torch-simulator ('torch') or per-group
PennyLane ('pennylane_pergroup') runs, e.g. to cross-check the simulators.

Outputs:
    results/analysis_summary.xlsx   — per-seed and aggregate tables
    results/plots/                  — bar charts + heatmaps
"""

import os
import sys
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')  # headless / no display needed
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches

try:
    import seaborn as sns
    HAS_SEABORN = True
except ImportError:
    HAS_SEABORN = False

_REPO_DIR = os.path.dirname(os.path.abspath(__file__))
if _REPO_DIR not in sys.path:
    sys.path.insert(0, _REPO_DIR)
from config.seeds import SEEDS  # noqa: E402  single pre-declared seed list

RESULTS_DIR = os.environ.get("QUAFFLE_RESULTS_DIR", os.path.join(_REPO_DIR, "results"))
# Where analysis_summary.xlsx and plots/ are written. Defaults to RESULTS_DIR;
# set QUAFFLE_ANALYSIS_OUT_DIR to analyse a read-only results tree.
ANALYSIS_OUT_DIR = os.environ.get("QUAFFLE_ANALYSIS_OUT_DIR", RESULTS_DIR)
DATASETS = ['optical', 'sar', 'combined']

# One pre-declared seed pool (config/seeds.py) for every method. Seeds are
# never dropped per method after looking at results. Missing or incomplete
# runs are reported by report_seed_coverage() and paired statistics use the
# per-cell intersection of completed seeds, always reporting n and the
# included and missing seeds.
MAIN_METHOD_SEEDS = set(SEEDS)
MAIN_METHOD_SEED_LABELS = {
    "Classical Central", "Quantum Central", "Ablation Central",
    "Classical Fed (Random)", "Quantum Fed (Random)", "Ablation Fed (Random)",
    "Classical Fed (Dirichlet)", "Quantum Fed (Dirichlet)", "Ablation Fed (Dirichlet)",
}

# ─────────────────────────────────────────────────────────────────────────────
# Publication style
# ─────────────────────────────────────────────────────────────────────────────

def _setup_pub_style():
    """Set matplotlib rcParams for clean, publication-quality figures."""
    plt.rcParams.update({
        # Fonts — serif (Times-like)
        'font.family':            'serif',
        'font.serif':             ['DejaVu Serif', 'Times New Roman', 'Times',
                                   'Liberation Serif', 'Computer Modern Roman'],
        'mathtext.fontset':       'dejavuserif',
        'font.size':              10,
        'axes.titlesize':         11,
        'axes.titleweight':       'bold',
        'axes.labelsize':         10,
        'xtick.labelsize':        8.5,
        'ytick.labelsize':        8.5,
        'legend.fontsize':        9,
        'legend.title_fontsize':  9,
        'figure.titlesize':       12,
        # Axes — kill top/right spines, thin lines
        'axes.spines.top':        False,
        'axes.spines.right':      False,
        'axes.linewidth':         0.8,
        'axes.edgecolor':         '#333333',
        'axes.labelcolor':        '#222222',
        'xtick.color':            '#333333',
        'ytick.color':            '#333333',
        'text.color':             '#222222',
        # Ticks
        'xtick.major.width':      0.8,
        'ytick.major.width':      0.8,
        'xtick.major.size':       3.0,
        'ytick.major.size':       3.0,
        'xtick.direction':        'out',
        'ytick.direction':        'out',
        # Grid (very subtle)
        'grid.linewidth':         0.5,
        'grid.alpha':             0.3,
        'grid.color':             '#888888',
        # Lines / patches
        'lines.linewidth':        1.5,
        'patch.linewidth':        0.4,
        'patch.edgecolor':        'white',
        # Legend
        'legend.frameon':         False,
        'legend.borderaxespad':   0.5,
        # Output
        'savefig.dpi':            300,
        'savefig.bbox':           'tight',
        'savefig.facecolor':      'white',
        'figure.facecolor':       'white',
        'figure.dpi':             100,
    })


_setup_pub_style()


# Structured colorblind-safe palette:
#   hue conveys method-type (Classical=blue, Quantum=orange/red, Ablation=green)
#   lightness conveys training regime (Central=dark, Fed Random=mid, Fed Dirichlet=light)
METHOD_PALETTE = {
    "Classical Central":         "#1F4E79",
    "Quantum Central":           "#B5482F",
    "Ablation Central":          "#2E6B40",
    "Classical Fed (Random)":    "#5B9BD5",
    "Quantum Fed (Random)":      "#ED7D31",
    "Ablation Fed (Random)":     "#70AD47",
    "Classical Fed (Dirichlet)": "#9DC3E6",
    "Quantum Fed (Dirichlet)":   "#F4B183",
    "Ablation Fed (Dirichlet)":  "#A9D18E",
}

# Palette echoes the PQC circuit figure: sky blue for the classical
# baseline (matches the RX/RY gate boxes), rose pink for the quantum/PQC method
# (matches the CNOT targets), warm gray for the parameter-matched ablation
# (a neutral third tone that complements both).
CATEGORY_COLOR = {
    "Classical": "#5A9AB8",
    "Quantum":   "#D4729E",
    "Ablation":  "#A89B8C",
}

DATASET_LABELS = {"optical": "Optical", "sar": "SAR", "combined": "Combined"}

# Ordered list of methods for consistent display (full — includes QC variants)
METHOD_ORDER = [
    "Classical Central",
    "Quantum Central",
    "Ablation Central",
    "Classical Fed (Random)",
    "Quantum Fed (Random)",
    "Quantum Fed (Random) QC=64",
    "Quantum Fed (Random) QC=128",
    "Quantum Fed (Random) QC=256",
    "Ablation Fed (Random)",
    "Classical Fed (Dirichlet)",
    "Quantum Fed (Dirichlet)",
    "Quantum Fed (Dirichlet) QC=64",
    "Quantum Fed (Dirichlet) QC=128",
    "Quantum Fed (Dirichlet) QC=256",
    "Ablation Fed (Dirichlet)",
]

# Methods used in main plots (bar, heatmap, seed lines, quantum-vs-classical).
# QUAFFLE is represented by the default QC=32 only — other QC values are
# reserved for the dedicated quantum-channels ablation plots.
MAIN_METHOD_ORDER = [
    "Classical Central",
    "Quantum Central",
    "Ablation Central",
    "Classical Fed (Random)",
    "Quantum Fed (Random)",
    "Ablation Fed (Random)",
    "Classical Fed (Dirichlet)",
    "Quantum Fed (Dirichlet)",
    "Ablation Fed (Dirichlet)",
]

METRIC_DISPLAY = {
    "Test_IoU":      "Test IoU",
    "Test_Dice":     "Test Dice",
    "Test_AUC":      "Test AUC",
    "Test_Accuracy": "Test Accuracy",
    "Val_IoU":       "Val IoU",
    "Val_Dice":      "Val Dice",
    "Val_Loss":      "Val Loss",
    "Test_Loss":     "Test Loss",
}

DATASET_COLORS = {"optical": "#4C72B0", "sar": "#DD8452", "combined": "#55A868"}

# ─────────────────────────────────────────────────────────────────────────────
# Data Loading
# ─────────────────────────────────────────────────────────────────────────────

def assign_method(row):
    """Return a clean method label from Training / Variant / Partition / Demo.

    Returns None for rows that should be dropped from plots (currently: the old
    base_channels=28 "classical-matched" ablation runs, superseded by the new
    bottleneck-matched ablation with demo_name containing "bottleneck").
    """
    training = str(row.get("Training", "")).lower()
    variant  = str(row.get("Variant",  "")).lower()
    partition= str(row.get("Partition","")).lower()
    demo     = str(row.get("Demo",     "")).lower()

    # "abalation" is the misspelled prefix written before the script was
    # renamed to ablation_studies.py. Accepted for backward compatibility.
    is_ablation = "abalation" in demo or "ablation" in demo
    # Drop the invalid old-architecture ablation rows (7.2M params, bc=28).
    # Keep only the new bottleneck-matched ablation rows (8.76M params).
    if is_ablation and "bottleneck" not in demo:
        return None

    # Detect non-default quantum_channels
    qc = row.get("Quantum_Channels", 32)
    try:
        qc = int(float(qc)) if pd.notna(qc) else 32
    except (ValueError, TypeError):
        qc = 32
    qc_suffix = f" QC={qc}" if ("quantum" in variant or is_ablation) and qc != 32 else ""

    if training == "centralized":
        if is_ablation:
            return f"Ablation Central{qc_suffix}"
        if "quantum" in variant:
            return f"Quantum Central{qc_suffix}"
        return "Classical Central"

    if training == "federated":
        if "dirichlet" in partition:
            part_label = "(Dirichlet)"
        elif "modality" in partition:
            part_label = "(Modality)"
        elif "geo_cluster" in partition or "geocluster" in partition:
            part_label = "(GeoCluster)"
        elif "random" in partition:
            part_label = "(Random)"
        else:
            part_label = ""

        if is_ablation:
            return f"Ablation Fed {part_label}{qc_suffix}".strip()
        if "quantum" in variant:
            return f"Quantum Fed {part_label}{qc_suffix}".strip()
        return f"Classical Fed {part_label}".strip()

    return "Unknown"


import re as _re_disk
import json as _json_disk

# Result directory naming convention:
#   flood_<central|federated>_optical_pytorch_<classical|pennylane>[_random|_dirichlet0.5][_combined|_sar][_qcN][<vtag>][<btag>]_seed<N>
#   ablation_<centralized|federated>_torch_bottleneck_matched_pennylane[_random|_dirichlet0.5][_combined|_sar][<vtag>]_seed<N>
#   (runs made before ablation_studies.py was renamed use the misspelled
#   prefix "abalation_", which is parsed the same way)
# <vtag> is the quantum ansatz variant, "_reupload_L<n>" or "_legacy"
# (quantum and ablation runs only, added 2026-10). Directories without a
# tag are older runs; their variant is read from the Provenance block of
# summary.txt (cli_args), see _infer_variant_from_summary.
# <btag> is the quantum circuit simulator (quantum runs only, added with
# --q-backend): '' for pennylane, "_qtorch" for the torch statevector
# simulator, "_qpergroup" for the per-group PennyLane loop. The q_backend
# recorded in metrics.json is authoritative; a metrics.json without the key
# was written before the q_backend option was recorded, when the model always
# used the torch simulator (LEGACY_Q_BACKEND).
# JAX dirs and old (non-bottleneck) ablation dirs are deliberately excluded.

# The variant whose Method labels stay unsuffixed ("Quantum Central", ...),
# so the summary tables and the forest plot keep working. Every other variant gets a
# " [<variant>]" suffix, e.g. "Quantum Fed (Random) [legacy]".
PRIMARY_VARIANT = os.environ.get("QUAFFLE_PRIMARY_VARIANT", "reupload_L2")
_DEFAULT_Q_LAYERS = 2

_SCALE_TAG_RE = _re_disk.compile(r"(?:_(?:bc\d+|img\d+|max\d+|dr[0-9a-f]{6}))+$")
_FED_MOD_PART = r"(?:_fedprox[\d.eE+-]+)?(?:_qagg-circular)?(?:_K\d+C\d+)?"
_CLASSICAL_MOD_PART = (
    r"(?:_(?:standard|depthwise_sep))?"
    r"(?:_in-(?:concat6|dual_stem))?"
    r"(?:_(?:dice_ce|focal_dice))?"
    r"(?:_notime)?(?:_aug)?"
    + _FED_MOD_PART
)
_VTAG_RE = _re_disk.compile(
    r"_((?:legacy|reupload_L\d+|hqunet_G\d+_L\d+)"
    r"(?:_nq\d+)?"
    r"(?:_(?:pixel|patch2x2|patch4x4|se_gate))?"
    r"(?:_abl-(?:linear|mlp|conv1x1|depthwise3x3|lowrank|hqunet_spatial(?:_G\d+)?(?:_L\d+)?))?"
    r"(?:_additive)?"
    + _CLASSICAL_MOD_PART
    + r")$"
)
_CLASSICAL_VTAG_RE = _re_disk.compile(
    r"_((?:"
    r"(?:standard|depthwise_sep|in-(?:concat6|dual_stem)|dice_ce|focal_dice|notime|aug|fedprox[\d.eE+-]+|K\d+C\d+)"
    r"(?:_(?:standard|depthwise_sep|in-(?:concat6|dual_stem)|dice_ce|focal_dice|notime|aug|fedprox[\d.eE+-]+|K\d+C\d+))*"
    r"))$"
)
_BTAG_RE = _re_disk.compile(r"_(qtorch|qpergroup)$")
_BTAG_TO_BACKEND = {"qtorch": "torch", "qpergroup": "pennylane_pergroup"}
Q_BACKENDS = ("pennylane", "pennylane_pergroup", "torch")
DEFAULT_INCLUDE_BACKENDS = ("pennylane",)
# metrics.json / summary.txt with provenance but no q_backend record:
# written before q_backend was recorded, the model used the torch simulator.
LEGACY_Q_BACKEND = "torch"
# summary.txt without a Provenance block: pre-revision code, which evaluated
# the circuit with per-item / per-group default.qubit QNodes (the torch
# simulator did not exist yet).
PRE_REVISION_Q_BACKEND = "pennylane_pergroup"
_BACKEND_LABEL = {"torch": "qtorch", "pennylane_pergroup": "qpergroup"}


def _backend_from_summary(text):
    """q_backend of a run that only has summary.txt."""
    m = _re_disk.search(r"Quantum circuit simulator:.*?q_backend=(\w+)", text or "")
    if m:
        return m.group(1)
    m = _re_disk.search(r"^\s*q_backend:\s*(\w+)", text or "", _re_disk.M)
    if m and m.group(1) in Q_BACKENDS:
        return m.group(1)
    if not _re_disk.search(r"cli_args:", text or ""):
        return PRE_REVISION_Q_BACKEND
    cli = _re_disk.search(r"cli_args:\s*(.*)", text).group(1)
    b = _re_disk.search(r"--q[-_]backend[ =](\w+)", cli)
    return b.group(1) if b else LEGACY_Q_BACKEND


def _with_backend(method, q_backend, label_backends):
    """Suffix ' [qtorch]' / ' [qpergroup]' when runs of several simulators are
    analysed together, so per-seed dedup never averages two simulators."""
    if not label_backends or q_backend in (None, "pennylane"):
        return method
    label = f" [{_BACKEND_LABEL.get(q_backend, q_backend)}]"
    if " QC=" in method:
        head, qc = method.split(" QC=", 1)
        return f"{head}{label} QC={qc}"
    return f"{method}{label}"


def _variant_label(variant):
    """Method-label suffix for a variant ('' for the primary variant)."""
    primary = os.environ.get("QUAFFLE_PRIMARY_VARIANT", PRIMARY_VARIANT)
    if not variant or variant == primary:
        return ""
    return f" [{variant.replace('_', ' ')}]"


def _variant_from_fields(q_ansatz, q_layers, data=None):
    parts = []
    if q_ansatz:
        if q_ansatz == "legacy":
            parts.append("legacy")
        elif q_ansatz == "hqunet":
            grid = int(data.get("hqunet_grid") or 2) if isinstance(data, dict) else 2
            layers = int(q_layers) if q_layers is not None else _DEFAULT_Q_LAYERS
            parts.append(f"hqunet_G{grid}_L{layers}")
        else:
            layers = int(q_layers) if q_layers is not None else _DEFAULT_Q_LAYERS
            parts.append(f"{q_ansatz}_L{layers}")
        if isinstance(data, dict):
            n_q = data.get("n_qubits")
            if q_ansatz in ("legacy", "reupload") and n_q is not None and int(n_q) != 4:
                parts.append(f"nq{int(n_q)}")
            q_sp = data.get("q_spatial_mode")
            abl = data.get("ablation_mode")
            if q_sp and q_sp != "global" and (q_ansatz != "hqunet" or abl in ("linear", "mlp")):
                parts.append(str(q_sp))
            if abl and abl != "auto":
                if abl == "hqunet_spatial" and q_ansatz != "hqunet":
                    grid = int(data.get("hqunet_grid") or 2)
                    if q_ansatz == "legacy":
                        layers_abl = int(q_layers) if q_layers is not None else _DEFAULT_Q_LAYERS
                        parts.append(f"abl-{abl}_G{grid}_L{layers_abl}")
                    else:
                        parts.append(f"abl-{abl}_G{grid}")
                else:
                    parts.append(f"abl-{abl}")
            skip_mod = data.get("skip_modulation")
            if skip_mod and skip_mod != "affine":
                parts.append(str(skip_mod))
    if isinstance(data, dict):
        bb = data.get("backbone_conv_mode")
        if bb and bb != "ws_conv":
            parts.append(str(bb))
        cim = data.get("combined_input_mode")
        if cim and cim != "shared3":
            parts.append(f"in-{cim}")
        loss_m = data.get("loss")
        if loss_m and loss_m != "ce":
            parts.append(str(loss_m))
        if data.get("use_time_emb") is False:
            parts.append("notime")
        if data.get("augment") is True:
            parts.append("aug")
        if data.get("training") == "federated":
            fl_strat = data.get("fl_strategy")
            mu = data.get("fedprox_mu")
            eff_fl = (
                fl_strat
                if fl_strat is not None
                else ("fedprox" if (mu is not None and float(mu) > 0.0) else "fedavg")
            )
            if eff_fl == "fedprox":
                mu_val = float(mu) if mu is not None else 0.01
                parts.append(f"fedprox{mu_val:g}")
            q_agg = data.get("q_angle_agg")
            if q_agg and q_agg != "linear":
                parts.append(f"qagg-{q_agg}")
            nc = data.get("num_clients")
            cpr = data.get("clients_per_round")
            if nc is not None:
                nc_i = int(nc)
                cpr_i = int(cpr) if cpr is not None else nc_i
                if nc_i != 4 or cpr_i != nc_i:
                    parts.append(f"K{nc_i}C{cpr_i}")
    return "_".join(parts) if parts else None


def _infer_variant_from_summary(text, quantumish=True):
    """Variant of an untagged (older) run directory.

    With a Provenance block the CLI flags decide (--q-ansatz default was
    'reupload', --q-layers default 2). Runs without provenance predate the
    revision and used the original (legacy) circuit."""
    m = _re_disk.search(r"cli_args:\s*(.*)", text or "")
    if not m:
        return "legacy" if quantumish else None
    cli = m.group(1)
    a = _re_disk.search(r"--q[-_]ansatz[ =](\w+)", cli)
    layers = _re_disk.search(r"--q[-_]layers[ =](\d+)", cli)
    n_qubits = _re_disk.search(r"--n[-_]qubits[ =](\d+)", cli)
    grid = _re_disk.search(r"--hqunet[-_]grid[ =](\d+)", cli)
    q_sp = _re_disk.search(r"--q[-_]spatial[-_]mode[ =](\w+)", cli)
    abl = _re_disk.search(r"--ablation[-_]mode[ =](\w+)", cli)
    skip_m = _re_disk.search(r"--skip[-_]modulation[ =](\w+)", cli)
    bb = _re_disk.search(r"--backbone[-_]conv[-_]mode[ =](\w+)", cli)
    loss_m = _re_disk.search(r"--loss[ =](\w+)", cli)
    no_time = bool(_re_disk.search(r"--no[-_]time[-_]emb\b", cli))
    aug = bool(_re_disk.search(r"--augment\b", cli))
    data = {
        "n_qubits": int(n_qubits.group(1)) if n_qubits else None,
        "hqunet_grid": int(grid.group(1)) if grid else None,
        "q_spatial_mode": q_sp.group(1) if q_sp else None,
        "ablation_mode": abl.group(1) if abl else None,
        "skip_modulation": skip_m.group(1) if skip_m else None,
        "backbone_conv_mode": bb.group(1) if bb else None,
        "loss": loss_m.group(1) if loss_m else None,
        "use_time_emb": False if no_time else True,
        "augment": True if aug else False,
    }
    q_ansatz = (a.group(1) if a else "reupload") if quantumish else None
    return _variant_from_fields(
        q_ansatz,
        int(layers.group(1)) if layers else None,
        data=data,
    )


def _parse_run_dir_is_mini(dirname):
    """True if dirname carries a non-default scale/data-root tag before _seed<N>."""
    m = _re_disk.match(r"^(.+)_seed(\d+)$", dirname)
    if not m:
        return False
    return bool(_SCALE_TAG_RE.search(m.group(1)))


def _parse_run_dir(dirname):
    """Map a result directory name → (Method, Dataset, Seed, Quantum_Channels,
    Training, Partition) or None if it should be ignored (JAX, old ablation,
    not a result dir, etc.).

    For quantum / ablation dirs without a variant tag the Method label is
    provisional; load_from_disk() appends the variant read from the run.
    See _parse_run_dir_full for the variant itself."""
    full = _parse_run_dir_full(dirname)
    return None if full is None else full[:6]


def _parse_run_dir_full(dirname):
    """Like _parse_run_dir plus a 7th element: the variant from the directory
    tag ('legacy', 'reupload_L2', ...), None if untagged or not applicable.
    The Method label here never carries the variant suffix.
    See _parse_run_dir_backend for the simulator tag."""
    m = _re_disk.match(r"^(.+)_seed(\d+)$", dirname)
    if not m:
        return None
    base, seed = m.group(1), int(m.group(2))

    if "jax" in base:
        return None  # JAX runs are not analysed (legacy code)

    sm = _SCALE_TAG_RE.search(base)
    if sm:
        base = base[:sm.start()]

    dir_backend = None
    bm = _BTAG_RE.search(base)
    if bm:
        dir_backend = _BTAG_TO_BACKEND[bm.group(1)]
        base = base[:bm.start()]

    variant = None
    vm = _VTAG_RE.search(base)
    if vm:
        variant = vm.group(1)
        base = base[:vm.start()]
    elif "classical" in base:
        cvm = _CLASSICAL_VTAG_RE.search(base)
        if cvm:
            variant = cvm.group(1)
            base = base[:cvm.start()]

    dataset = "combined" if "_combined" in base else ("sar" if "_sar" in base else "optical")

    qc_m = _re_disk.search(r"_qc(\d+)", base)
    qc = int(qc_m.group(1)) if qc_m else 32
    qc_suffix = f" QC={qc}" if qc != 32 else ""

    # "abalation_" is the misspelled prefix of runs made before the rename (backward compatible)
    if base.startswith("abalation_") or base.startswith("ablation_"):
        if "bottleneck" not in base:
            return None  # old base_channels=28 architecture; superseded
        if "centralized" in base:
            return (f"Ablation Central{qc_suffix}", dataset, seed, qc, "centralized", "centralized", variant)
        if "random" in base:
            return (f"Ablation Fed (Random){qc_suffix}", dataset, seed, qc, "federated", "random", variant)
        if "dirichlet" in base:
            return (f"Ablation Fed (Dirichlet){qc_suffix}", dataset, seed, qc, "federated", "dirichlet", variant)
        if "modality" in base:
            return (f"Ablation Fed (Modality){qc_suffix}", dataset, seed, qc, "federated", "modality", variant)
        if "geocluster" in base or "geo_cluster" in base:
            return (f"Ablation Fed (GeoCluster){qc_suffix}", dataset, seed, qc, "federated", "geo_cluster", variant)
        return None

    if not base.startswith("flood_"):
        return None

    if base.startswith("flood_central"):
        training, partition = "centralized", "centralized"
    elif base.startswith("flood_federated"):
        training = "federated"
        if "random" in base:
            partition = "random"
        elif "dirichlet" in base:
            partition = "dirichlet"
        elif "modality" in base:
            partition = "modality"
        elif "geocluster" in base or "geo_cluster" in base:
            partition = "geo_cluster"
        else:
            return None
    else:
        return None

    if "pennylane" in base:
        kind = "Quantum"
    elif "classical" in base:
        kind, qc, qc_suffix = "Classical", None, ""
    else:
        return None

    if training == "centralized":
        method = f"{kind} Central{qc_suffix}"
    elif partition == "random":
        method = f"{kind} Fed (Random){qc_suffix}"
    elif partition == "modality":
        method = f"{kind} Fed (Modality){qc_suffix}"
    elif partition == "geo_cluster":
        method = f"{kind} Fed (GeoCluster){qc_suffix}"
    else:
        method = f"{kind} Fed (Dirichlet){qc_suffix}"

    return (method, dataset, seed, qc if qc is not None else 32, training, partition, variant)


def _parse_run_dir_backend(dirname):
    """q_backend from the directory tag: 'torch' (_qtorch),
    'pennylane_pergroup' (_qpergroup) or None (no tag)."""
    m = _re_disk.match(r"^(.+)_seed(\d+)$", dirname)
    if not m:
        return None
    base = m.group(1)
    sm = _SCALE_TAG_RE.search(base)
    if sm:
        base = base[:sm.start()]
    bm = _BTAG_RE.search(base)
    return _BTAG_TO_BACKEND[bm.group(1)] if bm else None


def _with_variant(method, variant):
    """Insert the variant suffix before a ' QC=N' suffix so QC parsing
    (method.split('QC=')) keeps working."""
    label = _variant_label(variant)
    if not label:
        return method
    if " QC=" in method:
        head, qc = method.split(" QC=", 1)
        return f"{head}{label} QC={qc}"
    return f"{method}{label}"


_METRIC_KEYS = (("loss", "Loss"), ("accuracy", "Accuracy"), ("auc", "AUC"),
                ("iou", "IoU"), ("dice", "Dice"))


def _parse_metrics_json(path):
    """Parse a run's results/metrics.json (canonical output since 2026-10).
    Returns (metrics dict, info dict) or None if absent / incomplete."""
    if not os.path.exists(path):
        return None
    try:
        with open(path) as f:
            data = _json_disk.load(f)
    except (OSError, ValueError):
        return None
    test = data.get("test") or {}
    if test.get("iou") is None:
        return None
    out = {}
    for k, name in _METRIC_KEYS:
        if test.get(k) is not None:
            out[f"Test_{name}"] = float(test[k])
        v = (data.get("val") or {}).get(k)
        if v is not None:
            out[f"Val_{name}"] = float(v)
    if data.get("total_params") is not None:
        out["Total_Params"] = int(data["total_params"])
    is_mini = bool(data.get("is_mini_run"))
    if data.get("base_channels") not in (None, 32) or data.get("image_size") not in (None, 128):
        is_mini = True
    if data.get("max_samples") is not None:
        is_mini = True
    info = {
        "variant": _variant_from_fields(data.get("q_ansatz"), data.get("q_layers"), data=data),
        # Key absent: written before q_backend was recorded (torch simulator).
        "q_backend": data.get("q_backend", LEGACY_Q_BACKEND),
        "q_backend_recorded": "q_backend" in data,
        "is_mini_run": is_mini,
    }
    return out, info


def _parse_summary(path):
    """Parse a run's results/summary.txt. Returns dict with Test_/Val_ metrics
    and Total_Params, or None if the run is incomplete (no Test block)."""
    if not os.path.exists(path):
        return None
    try:
        with open(path) as f:
            text = f.read()
    except OSError:
        return None

    out = {}

    # Total parameters: main runs use "Total parameters: N", ablation summaries
    # use "Classical model params: N" (the actual ablation-model count).
    pm = (_re_disk.search(r"Total parameters:\s+([\d,]+)", text)
          or _re_disk.search(r"Classical model params:\s+([\d,]+)", text))
    if pm:
        out["Total_Params"] = int(pm.group(1).replace(",", ""))

    # Validation block — prefer "Best Validation [Results] (...):" if present,
    # falling back to "Final Validation [Results]:".
    val_m = (
        _re_disk.search(
            r"Best Validation(?: Results)?(?:\s*\([^)]*\))?:\s*\n((?:\s*\w+:\s+[\d.]+\s*\n)+)",
            text,
        )
        or _re_disk.search(
            r"Final Validation(?: Results)?:\s*\n((?:\s*\w+:\s+[\d.]+\s*\n)+)",
            text,
        )
    )
    if val_m:
        for line in val_m.group(1).splitlines():
            mm = _re_disk.match(r"\s*(\w+):\s+([\d.]+)", line)
            if mm:
                out[f"Val_{mm.group(1)}"] = float(mm.group(2))

    test_m = _re_disk.search(
        r"Test Results:\s*\n((?:\s*\w+:\s+[\d.]+\s*\n)+)", text)
    if test_m:
        for line in test_m.group(1).splitlines():
            mm = _re_disk.match(r"\s*(\w+):\s+([\d.]+)", line)
            if mm:
                out[f"Test_{mm.group(1)}"] = float(mm.group(2))

    if "Test_IoU" not in out:
        return None
    out["_variant_inferred"] = _infer_variant_from_summary(text, quantumish=True)
    out["_variant_inferred_classical"] = _infer_variant_from_summary(text, quantumish=False)
    out["_q_backend_inferred"] = _backend_from_summary(text)
    return out


def load_from_disk(label_backends=False, include_mini=None, results_dir=None):
    """Walk results/ and build the per-run DataFrame. Prefers each run's
    results/metrics.json (canonical, machine-readable) and falls back to
    summary.txt for older runs. Bypasses the Excel logging entirely.

    Returns a DataFrame with columns: Method, Dataset, Seed, Training,
    Partition, Quantum_Channels, Q_Variant, Source, Total_Params,
    Test_/Val_ {Loss, Accuracy, AUC, IoU, Dice}, plus Demo / Variant /
    Dirichlet_Alpha placeholders for backwards-compat with downstream code.
    Method labels carry a " [<variant>]" suffix for non-primary quantum
    ansatz variants (see PRIMARY_VARIANT). Q_Backend is the circuit
    simulator of quantum runs (None for classical and ablation runs); with
    label_backends=True non-pennylane quantum runs also get " [qtorch]" /
    " [qpergroup]".
    """
    if isinstance(label_backends, (str, os.PathLike)) and results_dir is None:
        results_dir = str(label_backends)
        label_backends = False
    rdir = results_dir or os.environ.get("QUAFFLE_RESULTS_DIR", RESULTS_DIR)
    if include_mini is None:
        include_mini = os.environ.get("QUAFFLE_INCLUDE_MINI", "0").lower() in ("1", "true", "yes")
    rows = []
    skipped_no_summary = 0
    skipped_mini = 0
    skipped_dirs = []
    n_json = 0
    if not os.path.isdir(rdir):
        return pd.DataFrame()

    for d in sorted(os.listdir(rdir)):
        full = os.path.join(rdir, d)
        if not os.path.isdir(full):
            continue
        parsed = _parse_run_dir_full(d)
        if parsed is None:
            continue
        method, dataset, seed, qc, training, partition, dir_variant = parsed
        quantumish = method.startswith("Quantum") or method.startswith("Ablation")
        dir_backend = _parse_run_dir_backend(d)
        is_mini = _parse_run_dir_is_mini(d)

        js = _parse_metrics_json(os.path.join(full, "results", "metrics.json"))
        if js is not None:
            metrics, info = js
            source = "metrics.json"
            variant = info.get("variant") or dir_variant
            q_backend = info["q_backend"]
            is_mini = is_mini or bool(info.get("is_mini_run"))
            if dir_backend and info["q_backend_recorded"] and dir_backend != q_backend:
                print(f"  WARNING: {d}: directory tag says q_backend={dir_backend!r}, "
                      f"metrics.json says {q_backend!r}; using metrics.json")
            elif dir_backend and not info["q_backend_recorded"]:
                q_backend = dir_backend
        else:
            metrics = _parse_summary(os.path.join(full, "results", "summary.txt"))
            if metrics is None:
                skipped_no_summary += 1
                skipped_dirs.append(d)
                continue
            source = "summary.txt"
            inferred = metrics.pop("_variant_inferred", None)
            inferred_cl = metrics.pop("_variant_inferred_classical", None)
            variant = dir_variant or (inferred if quantumish else inferred_cl)
            q_backend = dir_backend or metrics.pop("_q_backend_inferred", None)
            metrics.pop("_q_backend_inferred", None)
        if is_mini and not include_mini:
            skipped_mini += 1
            continue
        if source == "metrics.json":
            n_json += 1
        if not method.startswith("Quantum"):
            q_backend = None  # classical and ablation runs have no circuit
        rows.append({
            "Method": _with_backend(_with_variant(method, variant), q_backend, label_backends),
            "Dataset": dataset,
            "Seed": seed,
            "Training": training,
            "Partition": partition,
            "Quantum_Channels": qc,
            "Q_Variant": variant,
            "Q_Backend": q_backend,
            "Run_Dir": d,
            "Source": source,
            "Is_Mini": bool(is_mini),
            "Demo": method,                 # downstream filter compat
            "Variant": "quantum" if method.startswith("Quantum") else (
                       "ablation" if method.startswith("Ablation") else "classical"),
            "Dirichlet_Alpha": 0.5 if partition == "dirichlet" else None,
            **metrics,
        })

    if rows:
        print(f"  Parsed {len(rows)} runs ({n_json} from metrics.json, "
              f"{len(rows) - n_json} from summary.txt)")
    if skipped_mini:
        print(f"  Excluded {skipped_mini} mini/downscaled run dirs "
              f"(pass --include-mini or QUAFFLE_INCLUDE_MINI=1 to include)")
    if skipped_no_summary:
        print(f"  Skipped {skipped_no_summary} run dirs without test metrics "
              f"(failed, in-progress or crashed):")
        for d in skipped_dirs:
            print(f"    SKIPPED (no Test block): {d}")
    return pd.DataFrame(rows)


def filter_backends(df, include_backends=DEFAULT_INCLUDE_BACKENDS):
    """Keep classical / ablation rows and the quantum rows whose Q_Backend is
    in include_backends. Prints what was excluded."""
    if df.empty or "Q_Backend" not in df.columns:
        return df
    include = set(include_backends)
    quantum = df["Q_Backend"].notna()
    keep = ~quantum | df["Q_Backend"].isin(include)
    dropped = df[~keep]
    if len(dropped):
        counts = dropped["Q_Backend"].value_counts().to_dict()
        print(f"  Excluded {len(dropped)} quantum runs by circuit simulator {counts} "
              f"(included: {sorted(include)}; use --include-backend to change)")
    return df[keep].reset_index(drop=True)


def load_all(
    include_backends=DEFAULT_INCLUDE_BACKENDS,
    include_mini=None,
    results_dir=None,
    seeds=None,
):
    """Load all per-run results from disk (metrics.json and summary.txt files) and apply
    seed/method/simulator filters. Disk reads are authoritative — the Excel
    logging has historical column-shift bugs for some seeds (notably 7, 42,
    123, 256 in the optical sheet). Quantum runs are restricted to the
    circuit simulators in include_backends (default: pennylane only)."""
    print("Loading from disk (results/*/results/metrics.json and summary.txt)...")
    include_backends = tuple(include_backends)
    df = load_from_disk(
        label_backends=len(include_backends) > 1,
        include_mini=include_mini,
        results_dir=results_dir,
    )
    if df.empty:
        print("No data found.")
        return df
    q = df[df["Q_Backend"].notna()]
    if len(q):
        print(f"  Quantum runs by circuit simulator (q_backend): "
              f"{q['Q_Backend'].value_counts().to_dict()}")
    df = filter_backends(df, include_backends)
    print(f"  {len(df)} runs loaded across {df['Dataset'].nunique()} datasets, "
          f"{df['Method'].nunique()} methods")

    # Restrict every method to the single pre-declared seed list (or explicit --seeds override).
    seed_list = tuple(int(s) for s in seeds) if seeds is not None else tuple(SEEDS)
    allowed_seeds = set(seed_list)
    n_before = len(df)
    df = df[df["Seed"].isin(allowed_seeds)].reset_index(drop=True)
    if len(df) != n_before:
        print(f"  Dropped {n_before - len(df)} rows with seeds outside {list(seed_list)}")

    report_seed_coverage(df, seeds=seed_list)
    return df


def report_seed_coverage(df, methods=None, seeds=None):
    """Print, loudly, every (method, dataset) cell missing seeds from the
    pre-declared list. Returns {(method, dataset): sorted missing seeds}."""
    if methods is None:
        methods = sorted(MAIN_METHOD_SEED_LABELS)
    seed_list = tuple(int(s) for s in seeds) if seeds is not None else tuple(SEEDS)
    seed_set = set(seed_list)
    missing = {}
    for m in methods:
        for ds in DATASETS:
            have = set(int(x) for x in df[(df["Method"] == m) & (df["Dataset"] == ds)]["Seed"].dropna())
            miss = sorted(seed_set - have)
            if miss:
                missing[(m, ds)] = miss
    if missing:
        print("!" * 70)
        print(f"WARNING: missing completed runs for the pre-declared seed list {list(seed_list)}")
        for (m, ds), miss in sorted(missing.items()):
            print(f"  MISSING  {m:<28} {ds:<9} n={len(seed_list) - len(miss)}/{len(seed_list)}  seeds={miss}")
        print("!" * 70)
    else:
        print(f"  Seed coverage complete: every main cell has all {len(seed_list)} seeds")
    return missing


def paired_stats(deltas):
    """Mean, std, n, 95% t CI, paired t-test and Wilcoxon signed-rank p-values
    for a list of per-seed paired differences (numeric values or (seed, delta)
    2-tuples)."""
    from scipy import stats as _st
    vals = [
        item[1] if isinstance(item, (tuple, list)) and len(item) == 2 else item
        for item in deltas
    ]
    d = np.asarray(vals, dtype=float)
    n = int(d.size)
    out = {
        "n": n,
        "mean": None,
        "std": None,
        "ci95_half": None,
        "ci_low": None,
        "ci_high": None,
        "tcrit": None,
        "t_stat": None,
        "t_p": None,
        "wilcoxon_stat": None,
        "wilcoxon_p": None,
    }
    if n == 0:
        return out
    out["mean"] = float(d.mean())
    out["std"] = float(d.std(ddof=1)) if n > 1 else 0.0
    if n > 1:
        tc = float(_st.t.ppf(0.975, df=n - 1))
        out["tcrit"] = tc
        out["ci95_half"] = float(tc * out["std"] / np.sqrt(n))
        out["ci_low"] = float(out["mean"] - out["ci95_half"])
        out["ci_high"] = float(out["mean"] + out["ci95_half"])
        t = _st.ttest_1samp(d, 0.0)
        out["t_stat"], out["t_p"] = float(t.statistic), float(t.pvalue)
        if np.any(d != 0):
            try:
                w = _st.wilcoxon(d)
                out["wilcoxon_stat"], out["wilcoxon_p"] = float(w.statistic), float(w.pvalue)
            except ValueError as e:
                print(f"  Wilcoxon not computed: {e}")
    return out


# (regime, quantum method, classical comparator) pairs for paired tests.
PAIRED_COMPARISONS = [
    ("Centralized", "Quantum Central", "Classical Central"),
    ("Federated (Random)", "Quantum Fed (Random)", "Classical Fed (Random)"),
    ("Federated (Dirichlet)", "Quantum Fed (Dirichlet)", "Classical Fed (Dirichlet)"),
    ("Centralized ablation", "Quantum Central", "Ablation Central"),
    ("Federated (Random) ablation", "Quantum Fed (Random)", "Ablation Fed (Random)"),
    ("Federated (Dirichlet) ablation", "Quantum Fed (Dirichlet)", "Ablation Fed (Dirichlet)"),
]


def paired_tests(seed_df, metric="Test_IoU", seeds=None):
    """Paired quantum minus comparator tests per (regime, dataset) on the
    seeds both methods completed. Reports n, included and missing seeds,
    paired t-test and Wilcoxon signed-rank."""
    seed_set = set(int(s) for s in seeds) if seeds is not None else set(SEEDS)
    rows = []
    for regime, qm, cm in PAIRED_COMPARISONS:
        for ds in DATASETS:
            q = seed_df[(seed_df["Method"] == qm) & (seed_df["Dataset"] == ds)].set_index("Seed")[metric].dropna()
            c = seed_df[(seed_df["Method"] == cm) & (seed_df["Dataset"] == ds)].set_index("Seed")[metric].dropna()
            shared = sorted(int(x) for x in set(q.index) & set(c.index))
            if not shared and (q.empty or c.empty):
                continue
            deltas = [float(q.loc[s_] - c.loc[s_]) for s_ in shared]
            st = paired_stats(deltas)
            rows.append({
                "Regime": regime, "Dataset": ds, "Quantum": qm, "Comparator": cm,
                "Metric": metric, "N": st["n"],
                "Included_Seeds": ",".join(map(str, shared)),
                "Missing_Seeds": ",".join(map(str, sorted(seed_set - set(shared)))),
                "Delta_Mean": st["mean"], "Delta_Std": st["std"],
                "CI95_Half": st["ci95_half"],
                "T_Stat": st["t_stat"], "T_P": st["t_p"],
                "Wilcoxon_Stat": st["wilcoxon_stat"], "Wilcoxon_P": st["wilcoxon_p"],
            })
    return pd.DataFrame(rows)


# ─────────────────────────────────────────────────────────────────────────────
# Statistics
# ─────────────────────────────────────────────────────────────────────────────

def per_seed_summary(df):
    """
    Group by (Dataset, Method, Seed) and average metrics.
    This handles duplicate rows (e.g. if a job was resubmitted).
    """
    metrics = ["Test_IoU","Test_Dice","Test_AUC","Test_Accuracy",
               "Val_IoU","Val_Dice","Val_Loss","Test_Loss"]
    available = [m for m in metrics if m in df.columns]
    return (
        df.groupby(["Dataset","Method","Seed"])[available]
        .mean()
        .round(6)
        .reset_index()
    )


def aggregate_stats(seed_df, metric="Test_IoU"):
    """Mean ± std across seeds per (Dataset, Method)."""
    g = seed_df.groupby(["Dataset","Method"])[metric]
    agg = g.agg(["mean","std","count"]).reset_index()
    agg.columns = ["Dataset","Method","Mean","Std","N_Seeds"]
    agg["Mean±Std"] = agg.apply(
        lambda r: f"{r['Mean']:.4f} ± {r['Std']:.4f}" if pd.notna(r['Std'])
                  else f"{r['Mean']:.4f}", axis=1
    )
    return agg


def make_pivot(seed_df, metric="Test_IoU", method_order=None):
    """Pivot table: Method (rows) × Dataset (cols) with Mean±Std cells."""
    if method_order is None:
        method_order = METHOD_ORDER
    agg = aggregate_stats(seed_df, metric)
    pivot = agg.pivot(index="Method", columns="Dataset", values="Mean±Std")
    order = [m for m in method_order if m in pivot.index]
    leftover = [m for m in pivot.index if m not in order] if method_order is METHOD_ORDER else []
    pivot = pivot.reindex(order + leftover)
    ds_order = [d for d in DATASETS if d in pivot.columns]
    return pivot[ds_order]


def make_pivot_numeric(seed_df, metric="Test_IoU", method_order=None):
    """Same as make_pivot but with float Mean values (for heatmaps)."""
    if method_order is None:
        method_order = METHOD_ORDER
    agg = aggregate_stats(seed_df, metric)
    pivot = agg.pivot(index="Method", columns="Dataset", values="Mean")
    order = [m for m in method_order if m in pivot.index]
    leftover = [m for m in pivot.index if m not in order] if method_order is METHOD_ORDER else []
    pivot = pivot.reindex(order + leftover)
    ds_order = [d for d in DATASETS if d in pivot.columns]
    return pivot[ds_order].astype(float)


# ─────────────────────────────────────────────────────────────────────────────
# Excel Output
# ─────────────────────────────────────────────────────────────────────────────

def save_analysis_excel(df, seed_df, out_dir=None, seeds=None):
    target_dir = out_dir or os.environ.get("QUAFFLE_ANALYSIS_OUT_DIR", os.environ.get("QUAFFLE_RESULTS_DIR", ANALYSIS_OUT_DIR))
    os.makedirs(target_dir, exist_ok=True)
    out_path = os.path.join(target_dir, "analysis_summary.xlsx")

    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # ── Sheet 1: cleaned flat data ────────────────────────────────────────
        df.to_excel(writer, sheet_name="All_Data", index=False)

        # ── Sheet 2: per-seed summary ─────────────────────────────────────────
        seed_df.to_excel(writer, sheet_name="Per_Seed", index=False)

        # ── Sheets 3+: aggregate pivot per metric ─────────────────────────────
        for metric in ["Test_IoU","Test_Dice","Test_AUC","Test_Accuracy"]:
            pivot = make_pivot(seed_df, metric)
            pivot.to_excel(writer, sheet_name=f"Agg_{metric}")

        # ── Sheet: full aggregate (all metrics, long form) ────────────────────
        agg_rows = []
        for metric in ["Test_IoU","Test_Dice","Test_AUC","Test_Accuracy",
                       "Val_IoU","Val_Dice","Val_Loss","Test_Loss"]:
            if metric in seed_df.columns:
                a = aggregate_stats(seed_df, metric)
                a["Metric"] = metric
                agg_rows.append(a)
        pd.concat(agg_rows).to_excel(writer, sheet_name="Aggregate_Full", index=False)

        # ── Sheet: paired tests (t-test and Wilcoxon signed-rank) ─────────────
        pt = pd.concat([paired_tests(seed_df, m, seeds=seeds) for m in ("Test_IoU", "Test_Dice")
                        if m in seed_df.columns], ignore_index=True)
        pt.to_excel(writer, sheet_name="Paired_Tests", index=False)

    print(f"\nAnalysis Excel saved: {out_path}")


# ─────────────────────────────────────────────────────────────────────────────
# Plots
# ─────────────────────────────────────────────────────────────────────────────

def _method_colors(methods):
    """Color = method-type only (Classical/Ablation/Quantum). Same hue across regimes."""
    out = {}
    for m in methods:
        t = _method_type(m)
        out[m] = METHOD_TYPE_COLOR.get(t, "#666666")
    return out


def _annotate_bars(ax, bars, values, stds=None, fmt="{:.3f}", min_height_for_inside=None):
    """Annotate each bar with its value just above the error-bar top so the
    text doesn't collide with the error caps."""
    ymin, ymax = ax.get_ylim()
    pad = (ymax - ymin) * 0.018
    if stds is None:
        stds = [0.0] * len(values)
    for bar, v, s in zip(bars, values, stds):
        if v is None or (isinstance(v, float) and np.isnan(v)):
            continue
        s = 0.0 if (s is None or (isinstance(s, float) and np.isnan(s))) else s
        y = bar.get_height() + s + pad
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            y,
            fmt.format(v),
            ha="center", va="bottom",
            fontsize=7.5, color="#222222",
        )


# Order: Classical → Ablation → Quantum.  This puts the parameter-matched
# Ablation directly between the unrestricted Classical baseline and the Quantum
# model so the Quantum-vs-Ablation comparison is adjacent.
_CENTRALIZED_METHODS = ["Classical Central", "Ablation Central", "Quantum Central"]
_FEDERATED_METHODS   = ["Classical Fed (Random)",   "Ablation Fed (Random)",   "Quantum Fed (Random)",
                        "Classical Fed (Dirichlet)", "Ablation Fed (Dirichlet)", "Quantum Fed (Dirichlet)"]
_FED_RANDOM_METHODS    = ["Classical Fed (Random)",    "Ablation Fed (Random)",    "Quantum Fed (Random)"]
_FED_DIRICHLET_METHODS = ["Classical Fed (Dirichlet)", "Ablation Fed (Dirichlet)", "Quantum Fed (Dirichlet)"]

# Compact 3-method palette: hue conveys method-type, used uniformly across regimes.
# Echoes the PQC circuit palette: sky blue (gate boxes) for
# Classical, rose pink (CNOT targets) for Quantum, warm gray (neutral third)
# for the parameter-matched Ablation. Each (method, regime) tuple inherits the
# method's hue; regime is shown via subplot grouping, not color.
METHOD_TYPE_COLOR = {
    "Classical": "#5A9AB8",
    "Ablation":  "#A89B8C",
    "Quantum":   "#D4729E",
}

# Regime → hatch pattern (used only in combined 6+ method plots; primary plots
# don't hatch because they facet by regime).
REGIME_HATCH = {
    "Random":    "",
    "Dirichlet": "//",
}


def _method_type(method):
    if "Classical" in method: return "Classical"
    if "Ablation"  in method: return "Ablation"
    if "Quantum"   in method: return "Quantum"
    return "Other"


def _method_regime(method):
    if "Dirichlet" in method: return "Dirichlet"
    if "Random"    in method: return "Random"
    return "Centralized"


def _zoom_ylim(values, stds, padding=0.08, min_span=0.04):
    """Per-panel zoom: tight window around (value ± std) instead of [0, top]."""
    arr_v = np.asarray([v for v in values if not (v is None or (isinstance(v, float) and np.isnan(v)))])
    arr_s = np.asarray([s if not (s is None or (isinstance(s, float) and np.isnan(s))) else 0.0
                        for s in stds[:len(arr_v)]])
    if arr_v.size == 0:
        return (0.0, 1.0)
    lo = float(arr_v.min() - arr_s.max())
    hi = float(arr_v.max() + arr_s.max())
    span = hi - lo
    if span < min_span:
        mid = 0.5 * (hi + lo)
        lo, hi = mid - min_span / 2, mid + min_span / 2
        span = min_span
    pad = span * padding
    return (max(0.0, lo - pad), hi + pad)


def _bar_subplot(ax, sub, methods, metric, title, show_ylabel=True, annotate=False, ylim=None):
    """Bar chart for one dataset panel — zoomed y-axis, no headroom-padded ylim."""
    agg = aggregate_stats(sub, metric)
    agg = agg[agg["Method"].isin(methods)]
    agg = agg.set_index("Method").reindex(methods).reset_index()
    agg = agg.dropna(subset=["Mean"])
    if agg.empty:
        ax.set_title(f"{title}\n(no data)", pad=12)
        ax.set_xticks([]); ax.set_yticks([])
        for sp in ax.spines.values():
            sp.set_visible(False)
        return None

    palette = _method_colors(methods)
    colors  = [palette[m] for m in agg["Method"]]
    x = np.arange(len(agg))
    bars = ax.bar(
        x, agg["Mean"],
        yerr=agg["Std"].fillna(0).values,
        color=colors, edgecolor="white", linewidth=0.6,
        error_kw=dict(elinewidth=0.9, capsize=3.0, capthick=0.9, ecolor="#444444"),
    )
    ax.set_xticks(x)
    ax.set_xticklabels(agg["Method"], rotation=30, ha="right")
    ax.set_title(title, pad=10)
    if show_ylabel:
        ax.set_ylabel(METRIC_DISPLAY.get(metric, metric))
    ax.yaxis.grid(True, linestyle=":", alpha=0.5)
    ax.set_axisbelow(True)
    if ylim is not None:
        ax.set_ylim(*ylim)
    else:
        ax.set_ylim(*_zoom_ylim(agg["Mean"].tolist(), agg["Std"].fillna(0).tolist()))
    if annotate:
        _annotate_bars(ax, bars, agg["Mean"].tolist(),
                       stds=agg["Std"].fillna(0).tolist())
    return bars


def _dot_subplot(ax, sub, methods, metric, title, seed_df_full, show_ylabel=True,
                 show_individual=True, ylim=None):
    """Dot + 95%-CI plot, optionally with individual seed points as a strip."""
    agg = aggregate_stats(sub, metric)
    agg = agg[agg["Method"].isin(methods)]
    agg = agg.set_index("Method").reindex(methods).reset_index()
    agg = agg.dropna(subset=["Mean"])
    if agg.empty:
        ax.set_title(f"{title}\n(no data)", pad=12)
        ax.set_xticks([]); ax.set_yticks([])
        for sp in ax.spines.values():
            sp.set_visible(False)
        return None

    palette = _method_colors(methods)
    x = np.arange(len(agg))
    means = agg["Mean"].values.astype(float)
    stds  = agg["Std"].fillna(0).values.astype(float)
    n     = agg["N_Seeds"].fillna(1).values.astype(float)
    # 95% CI under the normal approximation (n is small, so this is loose)
    ci95 = 1.96 * stds / np.sqrt(np.clip(n, 1, None))

    # Faint background: each individual seed value as a small dot ("strip")
    if show_individual:
        for xi, m in zip(x, agg["Method"]):
            vals = sub[sub["Method"] == m][metric].dropna().values
            if len(vals) == 0:
                continue
            jitter = (np.random.default_rng(int(xi) + 7).uniform(-0.10, 0.10, size=len(vals)))
            ax.scatter(
                np.full_like(vals, xi, dtype=float) + jitter, vals,
                s=12, alpha=0.35, color=palette[m], edgecolors="none", zorder=2,
            )

    # Mean dot + 95% CI bar
    for xi, m, mean, hw in zip(x, agg["Method"], means, ci95):
        ax.errorbar(
            xi, mean, yerr=hw, fmt="o",
            color=palette[m], markerfacecolor=palette[m], markeredgecolor="white",
            markersize=7, linewidth=1.2, capsize=4, capthick=1.0, zorder=3,
        )

    ax.set_xticks(x)
    ax.set_xticklabels(agg["Method"], rotation=30, ha="right")
    ax.set_title(title, pad=10)
    if show_ylabel:
        ax.set_ylabel(METRIC_DISPLAY.get(metric, metric))
    ax.yaxis.grid(True, linestyle=":", alpha=0.5)
    ax.set_axisbelow(True)

    if ylim is not None:
        ax.set_ylim(*ylim)
    else:
        # Zoom to data range using individual values if available
        all_vals = []
        for m in agg["Method"]:
            all_vals.extend(sub[sub["Method"] == m][metric].dropna().tolist())
        if not all_vals:
            all_vals = means.tolist()
        lo, hi = min(all_vals), max(all_vals)
        span = hi - lo
        if span < 0.04:
            mid = 0.5 * (hi + lo); lo, hi = mid - 0.02, mid + 0.02; span = 0.04
        pad = span * 0.18
        ax.set_ylim(max(0.0, lo - pad), hi + pad)
    return None


_REGIME_METHOD_LISTS = {
    "centralized":   (_CENTRALIZED_METHODS,   "centralized training"),
    "fed_random":    (_FED_RANDOM_METHODS,    "IID FL (Random)"),
    "fed_dirichlet": (_FED_DIRICHLET_METHODS, "non-IID FL (Dirichlet α=0.5)"),
    "federated":     (_FEDERATED_METHODS,     "FL (Random + Dirichlet)"),
    "all":           (None,                   "all training regimes"),
}

_REGIME_TITLE = {
    "centralized":   "Parameter-matched comparison: centralized training",
    "fed_random":    "Parameter-matched FL comparison: random (IID) partitioning",
    "fed_dirichlet": "Parameter-matched FL comparison: Dirichlet (non-IID, α=0.5) partitioning",
    "federated":     "Parameter-matched FL comparison across partitioning regimes",
    "all":           "Parameter-matched comparison: all training regimes",
}


def _effective_n_seeds(seed_df, methods):
    """Floor of unique-seed counts across the (Dataset, Method) pairs that
    appear in a plot. Reflects the worst-covered cell, which is the honest n
    for cross-method comparisons (some cells can have fewer completed seeds).
    """
    if seed_df.empty or not methods:
        return 0
    counts = (
        seed_df[seed_df["Method"].isin(methods)]
        .groupby(["Dataset", "Method"])["Seed"].nunique()
    )
    return int(counts.min()) if not counts.empty else 0


def _ablation_caption_text():
    return ("Ablation: parameter-matched classical UNet (≈8.76M params, no PQC) — "
            "isolates the quantum contribution from raw model capacity.")


def _ablation_legend_handles():
    return [
        mpatches.Patch(color=METHOD_TYPE_COLOR["Classical"], label="Classical UNet"),
        mpatches.Patch(color=METHOD_TYPE_COLOR["Ablation"],
                       label="Ablation (parameter-matched classical, no PQC)"),
        mpatches.Patch(color=METHOD_TYPE_COLOR["Quantum"],   label="Quantum UNet"),
    ]


def plot_bar_comparison(seed_df, metric="Test_IoU", regime="centralized", style="bar", ylim=None):
    """Method × dataset comparison for one training regime.

    regime ∈ {"centralized", "fed_random", "fed_dirichlet", "federated", "all"}
    style  ∈ {"bar", "dot"} — "dot" plots mean ± 95% CI with individual seeds.
    ylim   — optional (lo, hi) tuple; if provided, used for every subplot so
             figures across regimes share a common Y axis.
    """
    methods, suffix = _REGIME_METHOD_LISTS.get(regime, (None, "All Methods"))
    if methods is None:
        methods = [m for m in MAIN_METHOD_ORDER if m in seed_df["Method"].unique()]
    methods = [m for m in methods if m in seed_df["Method"].unique()]
    if not methods:
        return None

    n_methods = len(methods)
    panel_w = 2.4 if n_methods <= 3 else (3.4 if n_methods <= 6 else 4.4)
    figsize = (panel_w * len(DATASETS), 3.7)

    fig, axes = plt.subplots(1, len(DATASETS), figsize=figsize, sharey=(ylim is not None))
    if len(DATASETS) == 1:
        axes = [axes]

    for ax, ds in zip(axes, DATASETS):
        sub = seed_df[seed_df["Dataset"] == ds]
        if style == "dot":
            _dot_subplot(ax, sub, methods, metric,
                         title=DATASET_LABELS[ds], seed_df_full=seed_df,
                         show_ylabel=(ax is axes[0]),
                         show_individual=True,
                         ylim=ylim)
        else:
            _bar_subplot(ax, sub, methods, metric,
                         title=DATASET_LABELS[ds],
                         show_ylabel=(ax is axes[0]),
                         annotate=True,
                         ylim=ylim)

    fig.legend(handles=_ablation_legend_handles(),
               loc="upper center", ncol=3, frameon=False,
               bbox_to_anchor=(0.5, 1.02))
    n_seeds = _effective_n_seeds(seed_df, methods)
    extra = "mean ± 95% CI, individual seeds shown" if style == "dot" else "mean ± std"
    title_prefix = _REGIME_TITLE.get(regime, "Parameter-matched comparison")
    fig.suptitle(
        f"{title_prefix}  (n≥{n_seeds} seeds, {extra})",
        y=1.10,
    )
    fig.tight_layout()
    return fig


def plot_seed_lines(seed_df, metric="Test_IoU", dataset="optical"):
    """Line chart: metric vs seed, one line per method, for a given dataset."""
    sub = seed_df[seed_df["Dataset"] == dataset].sort_values("Seed")
    if sub.empty:
        return None

    methods = [m for m in MAIN_METHOD_ORDER if m in sub["Method"].unique()]
    seeds = sorted(sub["Seed"].dropna().unique())
    fig, ax = plt.subplots(figsize=(7.0, 4.2))
    colors = _method_colors(methods)
    markers = ["o", "s", "^", "D", "v", "P", "X", "*", "h"]
    for i, method in enumerate(methods):
        m_data = sub[sub["Method"] == method].set_index("Seed")[metric].reindex(seeds)
        ax.plot(seeds, m_data.values, marker=markers[i % len(markers)],
                label=method, color=colors[method],
                linewidth=1.4, markersize=5, alpha=0.95)

    ax.set_xticks(seeds)
    ax.set_xticklabels([str(int(s)) for s in seeds], rotation=0)
    ax.set_xlabel("Seed")
    ax.set_ylabel(METRIC_DISPLAY.get(metric, metric))
    ax.set_title(f"{METRIC_DISPLAY.get(metric, metric)} per seed — {DATASET_LABELS[dataset]}")
    ax.legend(loc="center left", bbox_to_anchor=(1.02, 0.5), frameon=False)
    ax.yaxis.grid(True, linestyle=":", alpha=0.5)
    ax.set_axisbelow(True)
    fig.tight_layout()
    return fig


def plot_heatmap(seed_df, metric="Test_IoU"):
    """Heatmap of Method × Dataset means with cell annotations."""
    pivot = make_pivot_numeric(seed_df, metric, method_order=MAIN_METHOD_ORDER)
    if pivot.empty:
        return None
    pivot.columns = [DATASET_LABELS.get(c, str(c).capitalize()) for c in pivot.columns]

    fig, ax = plt.subplots(figsize=(max(4.0, len(pivot.columns) * 1.5),
                                    max(3.5, len(pivot) * 0.45 + 1.0)))
    cmap = "viridis" if metric.endswith("Loss") else "RdYlBu_r"
    if HAS_SEABORN:
        sns.heatmap(pivot, annot=True, fmt=".3f", cmap=cmap,
                    ax=ax, linewidths=0.6, linecolor="white",
                    cbar_kws={"shrink": 0.8, "label": METRIC_DISPLAY.get(metric, metric)},
                    annot_kws={"size": 8.5, "color": "#222222"})
    else:
        im = ax.imshow(pivot.values.astype(float), cmap=cmap, aspect="auto")
        cb = plt.colorbar(im, ax=ax, shrink=0.8)
        cb.set_label(METRIC_DISPLAY.get(metric, metric))
        for i in range(len(pivot)):
            for j in range(len(pivot.columns)):
                v = pivot.values[i, j]
                if not np.isnan(v):
                    ax.text(j, i, f"{v:.3f}", ha="center", va="center", fontsize=8.5)
        ax.set_xticks(range(len(pivot.columns)))
        ax.set_xticklabels(pivot.columns)
        ax.set_yticks(range(len(pivot)))
        ax.set_yticklabels(pivot.index)

    ax.set_title(f"{METRIC_DISPLAY.get(metric, metric)} — mean across seeds")
    ax.set_xlabel("")
    ax.set_ylabel("")
    fig.tight_layout()
    return fig


def plot_quantum_vs_classical(seed_df, metric="Test_IoU"):
    """Headline classical-vs-quantum comparison across training modes & datasets."""
    pairs = [
        ("Classical Central",         "Quantum Central",         "Centralized"),
        ("Classical Fed (Random)",    "Quantum Fed (Random)",    "Fed (Random)"),
        ("Classical Fed (Dirichlet)", "Quantum Fed (Dirichlet)", "Fed (Dirichlet)"),
    ]
    ds_list = [d for d in DATASETS if d in seed_df["Dataset"].unique()]
    n_ds = max(1, len(ds_list))
    # sharey=False: each panel zooms to its own data range so within-panel diffs are visible
    fig, axes = plt.subplots(1, n_ds, figsize=(3.6 * n_ds, 3.7), sharey=False)
    if n_ds == 1:
        axes = [axes]

    cls_color = CATEGORY_COLOR["Classical"]
    qtm_color = CATEGORY_COLOR["Quantum"]
    width = 0.36

    for ax, ds in zip(axes, ds_list):
        sub = seed_df[seed_df["Dataset"] == ds]
        x = np.arange(len(pairs))
        cls_m, cls_s, qtm_m, qtm_s = [], [], [], []
        for cls, qtm, _ in pairs:
            cv = sub[sub["Method"] == cls][metric]
            qv = sub[sub["Method"] == qtm][metric]
            cls_m.append(cv.mean() if len(cv) else np.nan)
            cls_s.append(cv.std() if len(cv) else 0.0)
            qtm_m.append(qv.mean() if len(qv) else np.nan)
            qtm_s.append(qv.std() if len(qv) else 0.0)

        b1 = ax.bar(x - width/2, cls_m, width,
                    yerr=np.nan_to_num(cls_s),
                    color=cls_color, edgecolor="white", linewidth=0.6,
                    error_kw=dict(elinewidth=0.9, capsize=2.5, capthick=0.9, ecolor="#444"),
                    label="Classical UNet")
        b2 = ax.bar(x + width/2, qtm_m, width,
                    yerr=np.nan_to_num(qtm_s),
                    color=qtm_color, edgecolor="white", linewidth=0.6,
                    error_kw=dict(elinewidth=0.9, capsize=2.5, capthick=0.9, ecolor="#444"),
                    label="Quantum UNet")

        ax.set_xticks(x)
        ax.set_xticklabels([p[2] for p in pairs], rotation=15, ha="right")
        ax.set_title(DATASET_LABELS[ds], pad=10)
        ax.yaxis.grid(True, linestyle=":", alpha=0.5)
        ax.set_axisbelow(True)
        if ax is axes[0]:
            ax.set_ylabel(METRIC_DISPLAY.get(metric, metric))
        # Per-panel zoom (no sharey) so within-panel diffs are visible
        all_v = [v for v in cls_m + qtm_m
                 if not (v is None or (isinstance(v, float) and np.isnan(v)))]
        all_s = [s for s in cls_s + qtm_s
                 if not (s is None or (isinstance(s, float) and np.isnan(s)))]
        ax.set_ylim(*_zoom_ylim(all_v, all_s))

    fig.legend(handles=[
        mpatches.Patch(color=cls_color, label="Classical UNet"),
        mpatches.Patch(color=qtm_color, label="Quantum UNet"),
    ], loc="upper center", bbox_to_anchor=(0.5, 1.02),
       ncol=2, frameon=False)
    cls_qtm_methods = [m for pair in pairs for m in pair[:2]]
    n_seeds = _effective_n_seeds(seed_df, cls_qtm_methods)
    fig.suptitle(
        f"Classical vs. quantum across training regimes  (n≥{n_seeds} seeds, mean ± std)",
        y=1.10,
    )
    fig.tight_layout()
    return fig


def plot_parameter_counts(df):
    """
    Horizontal bar chart of model parameter counts.
    One bar per model type, colour-coded by classical/quantum/ablation.
    """
    # Get one representative row per (Demo, Variant, Training) combo
    param_df = (
        df.dropna(subset=["Total_Params"])
        .groupby(["Demo", "Variant", "Training"])["Total_Params"]
        .first()
        .reset_index()
        .sort_values("Total_Params")
    )

    # Clean display names — short, consistent labels for plotting
    label_map = {
        "Centralized Classical Torch":                          "Classical\n(Centralized)",
        "Federated Classical Torch":                            "Classical\n(Federated)",
        "Centralized Quantum Torch":                            "Quantum\n(Centralized)",
        "Federated Quantum Torch":                              "Quantum\n(Federated)",
        # Old broken-arch (bc=28) ablations — should never appear post-filter, but
        # keep mapped in case any leak through. "Abalation" is the misspelled
        # demo name of runs made before the rename (backward compatible).
        "Abalation Centralized Torch Classical (vs pennylane)": "Ablation\n(Centralized)",
        "Abalation Federated Torch Classical (vs pennylane)":   "Ablation\n(Federated)",
        "Ablation Centralized Torch Classical":                 "Ablation\n(Centralized)",
        "Ablation Federated Torch Classical":                   "Ablation\n(Federated)",
        # New bottleneck-matched ablation (parameter-matched to ~0.01%)
        "Abalation Centralized Torch Bottleneck (vs pennylane)": "Ablation\n(Centralized)",
        "Abalation Federated Torch Bottleneck (vs pennylane)":   "Ablation\n(Federated)",
        "Ablation Centralized Torch Bottleneck":                 "Ablation\n(Centralized)",
        "Ablation Federated Torch Bottleneck":                   "Ablation\n(Federated)",
        "Ablation Centralized Torch Bottleneck (vs pennylane)":  "Ablation\n(Centralized)",
        "Ablation Federated Torch Bottleneck (vs pennylane)":    "Ablation\n(Federated)",
    }
    param_df["Label"] = param_df["Demo"].map(label_map).fillna(param_df["Demo"])

    # Deduplicate by label (same model appears in multiple datasets)
    param_df = param_df.drop_duplicates(subset=["Label"]).sort_values("Total_Params")

    # Colour by type
    def _color(row):
        d = row["Demo"].lower()
        if "abalation" in d or "ablation" in d:  # abalation: pre-rename runs
            return CATEGORY_COLOR["Ablation"]
        if "quantum" in row["Variant"]:
            return CATEGORY_COLOR["Quantum"]
        return CATEGORY_COLOR["Classical"]

    colors = param_df.apply(_color, axis=1).tolist()

    fig, ax = plt.subplots(figsize=(7.0, max(3.6, len(param_df) * 0.55)))
    bars = ax.barh(param_df["Label"], param_df["Total_Params"] / 1e6,
                   color=colors, edgecolor="white", height=0.6, linewidth=0.5)

    for bar, (_, row) in zip(bars, param_df.iterrows()):
        ax.text(bar.get_width() + ax.get_xlim()[1] * 0.012,
                bar.get_y() + bar.get_height() / 2,
                f"{row['Total_Params'] / 1e6:.2f}M",
                va="center", fontsize=8.5, color="#222222")

    ax.set_xlabel("Parameters (millions)")
    ax.set_title("Model parameter counts")
    ax.xaxis.grid(True, linestyle=":", alpha=0.5)
    ax.set_axisbelow(True)
    ax.set_xlim(right=ax.get_xlim()[1] * 1.18)

    legend_patches = [
        mpatches.Patch(color=CATEGORY_COLOR["Classical"], label="Classical UNet"),
        mpatches.Patch(color=CATEGORY_COLOR["Ablation"],  label="Ablation (parameter-matched)"),
        mpatches.Patch(color=CATEGORY_COLOR["Quantum"],   label="Quantum UNet"),
    ]
    # Place legend below the plot area so it can't cover bars
    fig.legend(handles=legend_patches, loc="lower center", ncol=3, frameon=False,
               bbox_to_anchor=(0.5, -0.04))
    fig.tight_layout()
    return fig


def _collect_param_table(df):
    """
    Build a tidy DataFrame keyed by (Training, Category, QC) with Total_Params.
    Category ∈ {Classical, Ablation, Quantum}. QC is int for Quantum, None otherwise.
    Filters Total_Params <= 1000 (corrupt rows logged as 0.0001), dedupes via median.
    """
    if "Total_Params" not in df.columns:
        return pd.DataFrame()
    sub = df.dropna(subset=["Total_Params"]).copy()
    sub["Total_Params"] = pd.to_numeric(sub["Total_Params"], errors="coerce")
    sub = sub[sub["Total_Params"] > 1000]
    if sub.empty:
        return pd.DataFrame()

    def _category(row):
        demo = str(row.get("Demo", "")).lower()
        variant = str(row.get("Variant", "")).lower()
        if "abalation" in demo or "ablation" in demo:  # abalation: pre-rename runs
            return "Ablation"
        if "quantum" in variant:
            return "Quantum"
        return "Classical"

    sub["Category"] = sub.apply(_category, axis=1)
    sub["Training"] = sub["Training"].astype(str).str.lower()

    def _qc(row):
        if row["Category"] != "Quantum":
            return None
        qc = row.get("Quantum_Channels", None)
        try:
            return int(float(qc)) if pd.notna(qc) else 32
        except (ValueError, TypeError):
            return 32

    sub["QC"] = sub.apply(_qc, axis=1)
    grouped = (
        sub.groupby(["Training", "Category", "QC"], dropna=False)["Total_Params"]
        .median()
        .reset_index()
    )
    return grouped


def _param_bar_data(df):
    """Return ordered list of (label, category, qc, params) per training mode."""
    tbl = _collect_param_table(df)
    if tbl.empty:
        return {}

    # Preferred ordering within each subplot
    ordering = [
        ("Classical", None),
        ("Ablation",  None),
        ("Quantum",   32),
        ("Quantum",   64),
        ("Quantum",   128),
        ("Quantum",   256),
    ]
    per_training = {}
    for training in ["centralized", "federated"]:
        sub = tbl[tbl["Training"] == training]
        rows = []
        for cat, qc in ordering:
            if qc is None:
                match = sub[(sub["Category"] == cat) & (sub["QC"].isna())]
            else:
                match = sub[(sub["Category"] == cat) & (sub["QC"] == qc)]
            if match.empty:
                continue
            label = {
                ("Classical", None): "Classical\nUNet",
                ("Ablation",  None): "Ablation\n(matched)",
            }.get((cat, qc), f"Quantum\nQC={qc}")
            params = float(match["Total_Params"].iloc[0])
            rows.append((label, cat, qc, params))
        per_training[training] = rows
    return per_training


def _bar_color(cat, qc):
    if cat == "Classical":
        return CATEGORY_COLOR["Classical"]
    if cat == "Ablation":
        return CATEGORY_COLOR["Ablation"]
    # Quantum — shade by QC value: alpha encodes channel count
    base_color = CATEGORY_COLOR["Quantum"]
    alpha_map = {32: 0.55, 64: 0.70, 128: 0.85, 256: 1.00}
    a = alpha_map.get(qc, 1.0)
    # Convert hex to rgba tuple
    r = int(base_color[1:3], 16) / 255
    g = int(base_color[3:5], 16) / 255
    b = int(base_color[5:7], 16) / 255
    return (r, g, b, a)


def _draw_param_bars(ax, rows, y_transform, ylabel, title, value_fmt="{:.2f}"):
    if not rows:
        ax.set_title(f"{title}\n(no data)")
        ax.set_xticks([]); ax.set_yticks([])
        for sp in ax.spines.values(): sp.set_visible(False)
        return
    labels = [r[0] for r in rows]
    values = [y_transform(r[3]) for r in rows]
    colors = [_bar_color(r[1], r[2]) for r in rows]
    x = np.arange(len(rows))
    bars = ax.bar(x, values, color=colors, edgecolor="white", linewidth=0.6, width=0.72)
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel(ylabel)
    ax.set_title(title, pad=10)
    ax.yaxis.grid(True, linestyle=":", alpha=0.5)
    ax.set_axisbelow(True)
    ax.set_ylim(0, max(values) * 1.28 if values else 1)
    pad = max(values) * 0.012 if values else 0.01
    for bar, v in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + pad,
                value_fmt.format(v), ha="center", va="bottom",
                fontsize=7.5, color="#222222")


def plot_parameter_bars(df):
    """Grouped bars: parameter counts by method type (centralized vs federated)."""
    data = _param_bar_data(df)
    if not data:
        return None
    fig, axes = plt.subplots(1, 2, figsize=(10.0, 4.0), sharey=True)
    _draw_param_bars(axes[0], data.get("centralized", []),
                     y_transform=lambda p: p / 1e6,
                     ylabel="Parameters (millions)",
                     title="Centralized",
                     value_fmt="{:.2f}M")
    _draw_param_bars(axes[1], data.get("federated", []),
                     y_transform=lambda p: p / 1e6,
                     ylabel="",
                     title="Federated",
                     value_fmt="{:.2f}M")
    legend_patches = [
        mpatches.Patch(color=CATEGORY_COLOR["Classical"], label="Classical UNet"),
        mpatches.Patch(color=CATEGORY_COLOR["Ablation"],  label="Ablation (param-matched)"),
        mpatches.Patch(color=CATEGORY_COLOR["Quantum"],   label="Quantum UNet (QC ∈ {32,64,128,256})"),
    ]
    fig.legend(handles=legend_patches, loc="upper center", ncol=3,
               bbox_to_anchor=(0.5, 1.04), frameon=False)
    fig.suptitle("Model parameter counts by method", y=1.10)
    fig.tight_layout()
    return fig


def plot_model_size_bars(df, bytes_per_param=4):
    """
    Grouped bar chart of model size in MB = params × bytes_per_param / 1e6.
    Default is float32 (4 bytes/param). Two subplots: Centralized vs Federated.
    Per-round FedAvg communication payload ≈ this size per client (up+down).
    """
    data = _param_bar_data(df)
    if not data:
        return None
    to_mb = lambda p: p * bytes_per_param / 1e6
    fig, axes = plt.subplots(1, 2, figsize=(10.0, 4.0), sharey=True)
    _draw_param_bars(axes[0], data.get("centralized", []),
                     y_transform=to_mb,
                     ylabel="Model size (MB)",
                     title="Centralized",
                     value_fmt="{:.1f} MB")
    _draw_param_bars(axes[1], data.get("federated", []),
                     y_transform=to_mb,
                     ylabel="",
                     title="Federated",
                     value_fmt="{:.1f} MB")
    legend_patches = [
        mpatches.Patch(color=CATEGORY_COLOR["Classical"], label="Classical UNet"),
        mpatches.Patch(color=CATEGORY_COLOR["Ablation"],  label="Ablation (param-matched)"),
        mpatches.Patch(color=CATEGORY_COLOR["Quantum"],   label="Quantum UNet (QC ∈ {32,64,128,256})"),
    ]
    fig.legend(handles=legend_patches, loc="upper center", ncol=3,
               bbox_to_anchor=(0.5, 1.04), frameon=False)
    fig.suptitle(f"Model size (float32, {bytes_per_param} B/param ≈ per-round FedAvg payload)",
                 y=1.10)
    fig.tight_layout()
    return fig


def plot_qc_ablation(seed_df, metric="Test_IoU"):
    """
    Line chart: quantum_channels on x-axis, metric on y-axis.
    One subplot per dataset, separate lines for Random vs Dirichlet.
    Includes classical baseline as a horizontal dashed line.
    """
    # Extract QC value from Method name
    qc_methods = seed_df[seed_df["Method"].str.startswith("Quantum Fed")].copy()
    if qc_methods.empty:
        return None

    def extract_qc(method):
        if "QC=" in method:
            return int(method.split("QC=")[1])
        return 32

    qc_methods["QC"] = qc_methods["Method"].apply(extract_qc)
    qc_methods["Partition"] = qc_methods["Method"].apply(
        lambda m: "Dirichlet" if "Dirichlet" in m else "Random"
    )

    ds_list = [d for d in DATASETS if d in qc_methods["Dataset"].unique()]
    if not ds_list:
        return None

    fig, axes = plt.subplots(1, len(ds_list), figsize=(4.0 * len(ds_list), 3.6), sharey=False)
    if len(ds_list) == 1:
        axes = [axes]

    partition_colors = {"Random": "#1F4E79", "Dirichlet": "#B5482F"}
    classical_styles = {"Random": dict(linestyle="--", color="#1F4E79"),
                        "Dirichlet": dict(linestyle=":",  color="#B5482F")}

    for ax, ds in zip(axes, ds_list):
        sub = qc_methods[qc_methods["Dataset"] == ds]
        for partition, marker in [("Random", "o"), ("Dirichlet", "s")]:
            p_sub = sub[sub["Partition"] == partition]
            if p_sub.empty:
                continue
            agg = p_sub.groupby("QC")[metric].agg(["mean", "std"]).reset_index().sort_values("QC")
            ax.errorbar(agg["QC"], agg["mean"], yerr=agg["std"].fillna(0),
                        marker=marker, capsize=2.5, label=f"Quantum ({partition})",
                        color=partition_colors[partition], linewidth=1.4, markersize=5,
                        elinewidth=0.9)
            cls_method = f"Classical Fed ({partition})"
            cls_vals = seed_df[(seed_df["Dataset"] == ds) & (seed_df["Method"] == cls_method)][metric]
            if not cls_vals.empty:
                ax.axhline(y=cls_vals.mean(),
                           **classical_styles[partition], alpha=0.55, linewidth=1.0,
                           label=f"Classical Fed ({partition})")

        ax.set_xlabel("Quantum channels")
        ax.set_xticks([32, 64, 128, 256])
        ax.set_xticklabels(["32", "64", "128", "256"])
        ax.set_title(DATASET_LABELS[ds])
        ax.yaxis.grid(True, linestyle=":", alpha=0.5)
        ax.set_axisbelow(True)
        if ax is axes[0]:
            ax.set_ylabel(METRIC_DISPLAY.get(metric, metric))
            ax.legend(loc="best", frameon=False)

    n_seeds = _effective_n_seeds(seed_df, qc_methods["Method"].unique().tolist())
    fig.suptitle(
        f"Quantum-channels ablation under FL  (n≥{n_seeds} seeds, mean ± std)",
        y=1.02,
    )
    fig.tight_layout()
    return fig


def plot_qc_ablation_bars(seed_df, metric="Test_IoU"):
    """
    Bar chart comparing different quantum_channels counts.
    One subplot per dataset; grouped bars: Random vs Dirichlet at each QC.
    Classical baseline shown as horizontal dashed line per partition.
    """
    qc_methods = seed_df[seed_df["Method"].str.startswith("Quantum Fed")].copy()
    if qc_methods.empty:
        return None

    def extract_qc(method):
        if "QC=" in method:
            return int(method.split("QC=")[1])
        return 32

    qc_methods["QC"] = qc_methods["Method"].apply(extract_qc)
    qc_methods["Partition"] = qc_methods["Method"].apply(
        lambda m: "Dirichlet" if "Dirichlet" in m else "Random"
    )

    ds_list = [d for d in DATASETS if d in qc_methods["Dataset"].unique()]
    if not ds_list:
        return None

    qc_values = sorted(qc_methods["QC"].unique())
    fig, axes = plt.subplots(1, len(ds_list), figsize=(4.0 * len(ds_list), 3.7), sharey=False)
    if len(ds_list) == 1:
        axes = [axes]

    partition_colors = {"Random": "#1F4E79", "Dirichlet": "#B5482F"}
    width = 0.38

    for ax, ds in zip(axes, ds_list):
        sub = qc_methods[qc_methods["Dataset"] == ds]
        x = np.arange(len(qc_values))
        panel_vals, panel_stds = [], []
        for i, partition in enumerate(["Random", "Dirichlet"]):
            p_sub = sub[sub["Partition"] == partition]
            if p_sub.empty:
                continue
            agg = p_sub.groupby("QC")[metric].agg(["mean", "std"]).reindex(qc_values).reset_index()
            offset = (i - 0.5) * width
            ax.bar(x + offset, agg["mean"].fillna(0).values, width,
                   yerr=agg["std"].fillna(0).values,
                   color=partition_colors[partition], edgecolor="white", linewidth=0.5,
                   error_kw=dict(elinewidth=0.9, capsize=2.5, capthick=0.9, ecolor="#444"),
                   label=f"Quantum ({partition})")
            panel_vals.extend(agg["mean"].dropna().tolist())
            panel_stds.extend(agg["std"].fillna(0).tolist())
            # Classical baseline (dashed line)
            cls_vals = seed_df[(seed_df["Dataset"] == ds)
                               & (seed_df["Method"] == f"Classical Fed ({partition})")][metric]
            if not cls_vals.empty:
                cls_mean = float(cls_vals.mean())
                ax.axhline(y=cls_mean, color=partition_colors[partition],
                           linestyle="--" if partition == "Random" else ":",
                           alpha=0.55, linewidth=1.0)
                panel_vals.append(cls_mean)
                panel_stds.append(0.0)

        ax.set_xticks(x)
        ax.set_xticklabels([str(q) for q in qc_values])
        ax.set_xlabel("Quantum channels")
        ax.set_title(DATASET_LABELS[ds], pad=10)
        ax.yaxis.grid(True, linestyle=":", alpha=0.5)
        ax.set_axisbelow(True)
        ax.set_ylim(*_zoom_ylim(panel_vals, panel_stds, padding=0.10))
        if ax is axes[0]:
            ax.set_ylabel(METRIC_DISPLAY.get(metric, metric))
            ax.legend(loc="best", frameon=False)

    n_seeds = _effective_n_seeds(seed_df, qc_methods["Method"].unique().tolist())
    fig.suptitle(
        f"Quantum-channels comparison under FL  (n≥{n_seeds} seeds; dashed line = classical baseline)",
        y=1.02,
    )
    fig.tight_layout()
    return fig


_REGIME_FOLDER_PREFIX = {
    "centralized":   "centralized",
    "fed_random":    "federated_random",
    "fed_dirichlet": "federated_dirichlet",
}

_REGIME_DISPLAY = {
    "centralized":   "Centralized",
    "fed_random":    "Federated (Random)",
    "fed_dirichlet": "Federated (Dirichlet α=0.5)",
}


def plot_bar_single_config(seed_df, metric, dataset, regime, ylim=None):
    """Bar plot for one (regime, dataset) cell — bars are the methods in that regime."""
    methods, _ = _REGIME_METHOD_LISTS.get(regime, (None, ""))
    if methods is None:
        return None
    methods = [m for m in methods if m in seed_df["Method"].unique()]
    if not methods:
        return None
    sub = seed_df[seed_df["Dataset"] == dataset]
    fig, ax = plt.subplots(1, 1, figsize=(4.5, 3.8))
    title = (f"{_REGIME_DISPLAY.get(regime, regime)} — "
             f"{DATASET_LABELS.get(dataset, dataset)} — "
             f"{METRIC_DISPLAY.get(metric, metric)}")
    _bar_subplot(ax, sub, methods, metric,
                 title=title, show_ylabel=True, annotate=True, ylim=ylim)
    fig.tight_layout()
    return fig


def _emit_per_config_bar_plots(seed_df, plots_dir):
    """Generate 9 (regime × dataset) folders × 4 metrics = 36 PNGs."""
    metrics = ["Test_Accuracy", "Test_AUC", "Test_IoU", "Test_Dice"]
    regimes = ["centralized", "fed_random", "fed_dirichlet"]

    # Compute a shared per-metric Y range across all 9 cells so the figures
    # for one metric are visually comparable across configs.
    def _shared_ylim_per_metric(metric):
        means, stds = [], []
        for reg in regimes:
            method_list, _ = _REGIME_METHOD_LISTS.get(reg, (None, ""))
            if method_list is None:
                continue
            method_list = [m for m in method_list if m in seed_df["Method"].unique()]
            for ds in DATASETS:
                sub = seed_df[seed_df["Dataset"] == ds]
                agg = aggregate_stats(sub, metric)
                agg = agg[agg["Method"].isin(method_list)].dropna(subset=["Mean"])
                means.extend(agg["Mean"].tolist())
                stds.extend(agg["Std"].fillna(0).tolist())
        if not means:
            return None
        return _zoom_ylim(means, stds)

    metric_ylim = {m: _shared_ylim_per_metric(m) for m in metrics}

    for reg in regimes:
        prefix = _REGIME_FOLDER_PREFIX[reg]
        for ds in DATASETS:
            folder = os.path.join(plots_dir, f"{prefix}_{ds}")
            os.makedirs(folder, exist_ok=True)
            for metric in metrics:
                fig = plot_bar_single_config(seed_df, metric, ds, reg,
                                             ylim=metric_ylim[metric])
                _save(fig, os.path.join(folder, f"{metric}.png"))


def _save(fig, path, dpi=300):
    if fig is None:
        return
    fig.savefig(path, dpi=dpi, bbox_inches="tight")
    plt.close(fig)
    print(f"  Saved {path}")


def save_all_plots(seed_df, raw_df=None, out_dir=None):
    """Output policy:
        results/plots/combinedcharts/                — legacy 3-in-1 figures (all metrics)
        results/plots/combinedcharts/supplementary/  — supp variants (Acc + AUC dot/seed lines)
        results/plots/<regime>_<dataset>/            — per-config bar plot (one PNG per metric)
    """
    target_dir = out_dir or os.environ.get("QUAFFLE_ANALYSIS_OUT_DIR", os.environ.get("QUAFFLE_RESULTS_DIR", ANALYSIS_OUT_DIR))
    plots_dir = os.path.join(target_dir, "plots")
    combined_dir = os.path.join(plots_dir, "combinedcharts")
    supp_dir  = os.path.join(combined_dir, "supplementary")
    os.makedirs(plots_dir, exist_ok=True)
    os.makedirs(combined_dir, exist_ok=True)
    os.makedirs(supp_dir, exist_ok=True)
    src = raw_df if raw_df is not None else seed_df

    # Model-architecture summaries (independent of metric)
    _save(plot_parameter_counts(src), os.path.join(combined_dir, "parameter_counts.png"))
    _save(plot_parameter_bars(src),    os.path.join(combined_dir, "parameter_bars_by_method.png"))
    _save(plot_model_size_bars(src),   os.path.join(combined_dir, "model_size_bars_by_method.png"))

    PRIMARY_METRICS = ["Test_IoU", "Test_Dice"]
    SUPP_METRICS    = ["Test_Accuracy", "Test_AUC"]

    def _shared_ylim_for(metric, regimes, style):
        """Common Y range across the per-regime plots for one metric, so the
        centralized / fed_random / fed_dirichlet (and combined federated) figures
        all use the same scale and are visually comparable."""
        means, stds, raw_vals = [], [], []
        for reg in regimes:
            methods, _ = _REGIME_METHOD_LISTS.get(reg, (None, ""))
            if methods is None:
                continue
            methods = [m for m in methods if m in seed_df["Method"].unique()]
            for ds in DATASETS:
                sub = seed_df[seed_df["Dataset"] == ds]
                agg = aggregate_stats(sub, metric)
                agg = agg[agg["Method"].isin(methods)].dropna(subset=["Mean"])
                means.extend(agg["Mean"].tolist())
                stds.extend(agg["Std"].fillna(0).tolist())
                if style == "dot":
                    for m in methods:
                        raw_vals.extend(sub[sub["Method"] == m][metric].dropna().tolist())
        if style == "dot" and raw_vals:
            return _zoom_ylim(raw_vals, [0.0] * len(raw_vals), padding=0.18)
        if not means:
            return None
        return _zoom_ylim(means, stds)

    def _emit_metric(metric, base_dir):
        """Generate the per-regime bar/dot variants and headline plots for one metric."""
        bar_regimes = ("centralized", "fed_random", "fed_dirichlet", "federated")
        dot_regimes = ("centralized", "fed_random", "fed_dirichlet")
        bar_ylim = _shared_ylim_for(metric, bar_regimes, "bar")
        dot_ylim = _shared_ylim_for(metric, dot_regimes, "dot")

        # Per-regime headline figures — 3 methods (Classical, Ablation, Quantum)
        for reg in ("centralized", "fed_random", "fed_dirichlet"):
            _save(plot_bar_comparison(seed_df, metric, regime=reg, style="bar", ylim=bar_ylim),
                  os.path.join(base_dir, f"bar_{metric}_{reg}.png"))
            _save(plot_bar_comparison(seed_df, metric, regime=reg, style="dot", ylim=dot_ylim),
                  os.path.join(base_dir, f"dot_{metric}_{reg}.png"))
        # Combined federated (6 methods, partition adjacent)
        _save(plot_bar_comparison(seed_df, metric, regime="federated", style="bar", ylim=bar_ylim),
              os.path.join(base_dir, f"bar_{metric}_federated.png"))
        # Heatmap & classical-vs-quantum headline
        _save(plot_heatmap(seed_df, metric),
              os.path.join(base_dir, f"heatmap_{metric}.png"))
        _save(plot_quantum_vs_classical(seed_df, metric),
              os.path.join(base_dir, f"quantum_vs_classical_{metric}.png"))
        # Quantum-channels ablation
        _save(plot_qc_ablation(seed_df, metric),
              os.path.join(base_dir, f"qc_ablation_{metric}.png"))
        _save(plot_qc_ablation_bars(seed_df, metric),
              os.path.join(base_dir, f"qc_ablation_bars_{metric}.png"))

    for metric in PRIMARY_METRICS:
        _emit_metric(metric, combined_dir)
    for metric in SUPP_METRICS:
        _emit_metric(metric, supp_dir)

    # Per-seed line charts (variability check) — primary for IoU+Dice, supp for the rest
    for ds in DATASETS:
        for metric in PRIMARY_METRICS:
            _save(plot_seed_lines(seed_df, metric, ds),
                  os.path.join(combined_dir, f"seed_lines_{ds}_{metric}.png"))
        for metric in SUPP_METRICS:
            _save(plot_seed_lines(seed_df, metric, ds),
                  os.path.join(supp_dir,  f"seed_lines_{ds}_{metric}.png"))

    # Per-config bar plots: one folder per (regime × dataset), 4 PNGs each
    # (Test_Accuracy, Test_AUC, Test_IoU, Test_Dice). Bars show Classical /
    # Quantum / Ablation methods within that regime.
    _emit_per_config_bar_plots(seed_df, plots_dir)


# ─────────────────────────────────────────────────────────────────────────────
# Console Summary
# ─────────────────────────────────────────────────────────────────────────────

def print_summary(seed_df, seeds=None):
    print("\n" + "="*70)
    print("AGGREGATE RESULTS: Test IoU (mean ± std across seeds)")
    print("="*70)
    pivot = make_pivot(seed_df, "Test_IoU")
    print(pivot.to_string())

    print("\n" + "="*70)
    print("AGGREGATE RESULTS: Test Dice (mean ± std across seeds)")
    print("="*70)
    pivot = make_pivot(seed_df, "Test_Dice")
    print(pivot.to_string())

    print("\n" + "="*70)
    print("SEED COUNTS per (Dataset, Method)")
    print("="*70)
    counts = seed_df.groupby(["Dataset","Method"])["Seed"].nunique().unstack("Dataset")
    print(counts.to_string())
    print()

    for metric in ("Test_IoU", "Test_Dice"):
        if metric not in seed_df.columns:
            continue
        print("=" * 70)
        print(f"PAIRED TESTS, {metric}: quantum minus comparator, paired by seed "
              f"(t-test and Wilcoxon signed-rank)")
        print("=" * 70)
        pt = paired_tests(seed_df, metric, seeds=seeds)
        for _, r in pt.iterrows():
            if r["N"] == 0:
                print(f"  {r['Regime']:<32}{r['Dataset']:<9} n=0 (no shared seeds)")
                continue
            tp = "NA" if r["T_P"] is None or pd.isna(r["T_P"]) else f"{r['T_P']:.4f}"
            wp = "NA" if r["Wilcoxon_P"] is None or pd.isna(r["Wilcoxon_P"]) else f"{r['Wilcoxon_P']:.4f}"
            miss = f"  missing={r['Missing_Seeds']}" if r["Missing_Seeds"] else ""
            print(f"  {r['Regime']:<32}{r['Dataset']:<9} n={r['N']:<3} delta={r['Delta_Mean']:+.4f} "
                  f"t_p={tp} wilcoxon_p={wp}{miss}")
        print()


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def parse_args(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description="Analyze QUAFFLE experiment results")
    ap.add_argument("--include-backend", nargs="+", default=list(DEFAULT_INCLUDE_BACKENDS),
                    choices=list(Q_BACKENDS) + ["all"], metavar="Q_BACKEND",
                    help="Circuit simulators (q_backend) of quantum runs to include: "
                         f"{', '.join(Q_BACKENDS)} or all (default: pennylane). "
                         "With more than one, non-pennylane methods get a "
                         "' [qtorch]' / ' [qpergroup]' label suffix.")
    ap.add_argument("--include-mini", action="store_true",
                    help="Include mini / downscaled verification runs in analysis")
    ap.add_argument("--results-dir", default=None,
                    help="Override input results directory (default: $QUAFFLE_RESULTS_DIR or results/)")
    ap.add_argument("--out-dir", default=None,
                    help="Override output directory for analysis_summary.xlsx and plots/")
    ap.add_argument("--skip-plots", action="store_true",
                    help="Skip generating PNG plots (write analysis_summary.xlsx only)")
    ap.add_argument("--seeds", type=int, nargs="+", default=None,
                    help="Override canonical seed list for filtering and coverage checks")
    args = ap.parse_args(argv)
    if "all" in args.include_backend:
        args.include_backend = list(Q_BACKENDS)
    return args


def main(argv=None):
    # Parse CLI arguments (including --results-dir, --out-dir, --include-mini, --skip-plots, --seeds)
    _args = parse_args(argv)
    print("Loading data...")
    inc_mini = True if _args.include_mini else None
    df = load_all(
        include_backends=_args.include_backend,
        include_mini=inc_mini,
        results_dir=_args.results_dir,
        seeds=_args.seeds,
    )

    if df.empty:
        print("No data found, make sure results/*/results/metrics.json or summary.txt exist.")
        sys.exit(1)

    print(f"\nTotal rows: {len(df)}")
    print(f"Datasets:   {sorted(df['Dataset'].unique())}")
    print(f"Methods:    {sorted(df['Method'].unique())}")
    print(f"Seeds:      {sorted(df['Seed'].dropna().unique().astype(int))}")

    # Deduplicate: average over exact duplicate (dataset, method, seed) combos
    seed_df = per_seed_summary(df)
    print(f"Per-seed rows after dedup: {len(seed_df)}")

    print_summary(seed_df, seeds=_args.seeds)

    out_dir = _args.out_dir or _args.results_dir
    print("\nSaving Excel...")
    save_analysis_excel(df, seed_df, out_dir=out_dir, seeds=_args.seeds)

    skip_plots = _args.skip_plots or (
        os.environ.get("QUAFFLE_SKIP_ANALYSIS_PLOTS", "0").lower() in ("1", "true", "yes")
    )
    if not skip_plots:
        print("\nGenerating plots...")
        save_all_plots(seed_df, raw_df=df, out_dir=out_dir)

    print("\nDone.")


if __name__ == "__main__":
    main()
