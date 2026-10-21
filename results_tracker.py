#!/usr/bin/env python3
"""
Shared utility to append run results to a central Excel summary file.
Each demo calls save_results_to_excel() at the end of training.
Results are split into per-dataset files: results_summary_optical.xlsx,
results_summary_sar.xlsx, results_summary_combined.xlsx.
"""

import os
from datetime import datetime


_RESULTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results")
# Testing / relocation hook: QUAFFLE_TRACKER_DIR overrides where the shared
# workbooks live.
_RESULTS_DIR = os.environ.get("QUAFFLE_TRACKER_DIR", _RESULTS_DIR)

COLUMNS = [
    "Timestamp",
    "Demo",
    "Framework",
    "Training",
    "Variant",
    "Dataset",
    # Validation metrics
    "Val_Loss",
    "Val_Accuracy",
    "Val_AUC",
    "Val_IoU",
    "Val_Dice",
    # Test metrics
    "Test_Loss",
    "Test_Accuracy",
    "Test_AUC",
    "Test_IoU",
    "Test_Dice",
    # Config
    "Partition",
    "Dirichlet_Alpha",
    "Epochs_Steps_Rounds",
    "Num_Clients",
    "Local_Epochs",
    "Base_Channels",
    "Quantum_Channels",
    "Batch_Size",
    "Learning_Rate",
    "Total_Params",
    # Provenance (tools/provenance.py). Older sheets get these header cells
    # appended the first time a new row is written.
    "Git_SHA",
    "Git_Dirty",
    "Git_Branch",
    "CLI_Args",
    "Seed",
    "Hostname",
    "SLURM_Job_ID",
    "Torch_Version",
    "PennyLane_Version",
    "Flwr_Version",
]

_PROVENANCE_COLUMN_KEYS = [
    ("Git_SHA", "git_sha"),
    ("Git_Dirty", "git_dirty"),
    ("Git_Branch", "git_branch"),
    ("CLI_Args", "cli_args"),
    ("Seed", "seed"),
    ("Hostname", "hostname"),
    ("SLURM_Job_ID", "slurm_job_id"),
    ("Torch_Version", "torch_version"),
    ("PennyLane_Version", "pennylane_version"),
    ("Flwr_Version", "flwr_version"),
]


def save_results_to_excel(demo_name, framework, training, variant, val_metrics, test_metrics, config,
                          dataset='optical', provenance=None):
    """
    Append one row of results to results/results_summary_{dataset}.xlsx.

    Args:
        demo_name:    Human-readable name, e.g. "Centralized Classical Torch"
        framework:    "torch" | "jax"
        training:     "centralized" | "federated"
        variant:      "classical" | "quantum"
        val_metrics:  dict with loss/accuracy/auc/iou/dice, or None
        test_metrics: dict with loss/accuracy/auc/iou/dice, or None
        config:       dict or ConfigDict with training hyperparameters
        dataset:      "optical" | "sar" | "combined"
        provenance:   dict from tools.provenance.collect_provenance(), or None

    Rows are only written for completed runs. If test_metrics is missing the
    run is treated as failed and nothing is appended.
    """
    if not test_metrics:
        print("Not writing a tracker row: no test metrics (failed or incomplete run).")
        return
    try:
        import openpyxl
    except ImportError:
        print("Warning: openpyxl not installed, skipping Excel export. Install with: pip install openpyxl")
        return

    excel_path = os.path.join(_RESULTS_DIR, f"results_summary_{dataset}.xlsx")
    os.makedirs(_RESULTS_DIR, exist_ok=True)

    # Cross-node lock. fcntl.flock is not reliable across nodes on GPFS (on
    # 2026-10-02 two jobs on different nodes wrote results_summary_optical.xlsx
    # at the same time and one crashed with BadZipFile). POSIX record locks
    # (fcntl.lockf) are honoured cluster-wide by GPFS. An O_CREAT|O_EXCL marker
    # file with a stale timeout is taken as well, as a second
    # guard for filesystems where record locks are node-local.
    with _ExcelLock(excel_path + ".lock"):
        _save_locked(excel_path, demo_name, framework, training, variant, dataset,
                     val_metrics, test_metrics, config, provenance)


_LOCK_TIMEOUT_S = 300   # give up waiting after this long
_STALE_LOCK_S = 600     # an O_EXCL lock older than this is treated as stale


class _ExcelLock:
    """fcntl.lockf (POSIX record lock) plus an O_CREAT|O_EXCL marker file.

    Raises TimeoutError if the lock cannot be taken within _LOCK_TIMEOUT_S;
    callers wrap the tracker so that never fails a finished run."""

    def __init__(self, lock_path):
        self.lock_path = lock_path
        self.excl_path = lock_path + ".excl"
        self.fd = None
        self.have_excl = False

    def __enter__(self):
        import fcntl
        import socket
        import time
        deadline = time.time() + _LOCK_TIMEOUT_S
        self.fd = os.open(self.lock_path, os.O_RDWR | os.O_CREAT, 0o664)
        while True:
            try:
                fcntl.lockf(self.fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except OSError:
                if time.time() > deadline:
                    os.close(self.fd)
                    self.fd = None
                    raise TimeoutError(f"could not lock {self.lock_path} in {_LOCK_TIMEOUT_S}s")
                time.sleep(1.0)
        owner = f"{socket.gethostname()} pid={os.getpid()} {datetime.now().isoformat()}\n"
        while True:
            try:
                xfd = os.open(self.excl_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o664)
                os.write(xfd, owner.encode())
                os.close(xfd)
                self.have_excl = True
                break
            except FileExistsError:
                try:
                    age = time.time() - os.path.getmtime(self.excl_path)
                except OSError:
                    continue  # vanished between the calls, retry now
                if age > _STALE_LOCK_S:
                    print(f"Warning: removing stale tracker lock {self.excl_path} "
                          f"({age:.0f}s old)")
                    try:
                        os.remove(self.excl_path)
                    except OSError:
                        pass
                    continue
                if time.time() > deadline:
                    self.__exit__(None, None, None)
                    raise TimeoutError(f"tracker lock {self.excl_path} held for too long")
                time.sleep(1.0)
        return self

    def __exit__(self, *exc):
        import fcntl
        if self.have_excl:
            try:
                os.remove(self.excl_path)
            except OSError:
                pass
            self.have_excl = False
        if self.fd is not None:
            try:
                fcntl.lockf(self.fd, fcntl.LOCK_UN)
            finally:
                os.close(self.fd)
                self.fd = None
        return False


def _load_or_new_workbook(excel_path):
    """Load the workbook; a corrupt file is moved aside (kept) and a new
    workbook is started instead of crashing."""
    import openpyxl
    if os.path.exists(excel_path):
        try:
            return openpyxl.load_workbook(excel_path)
        except Exception as e:
            dst = f"{excel_path}.corrupt-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{os.getpid()}"
            print(f"Warning: {excel_path} is unreadable ({type(e).__name__}: {e}). "
                  f"Moving it to {dst} and starting a new workbook. Per-run "
                  f"results/metrics.json files hold every run's metrics.")
            os.replace(excel_path, dst)
    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    return wb


def _atomic_save_workbook(wb, excel_path):
    import socket
    tmp = f"{excel_path}.tmp.{socket.gethostname()}.{os.getpid()}.xlsx"
    try:
        wb.save(tmp)
        with open(tmp, "rb") as f:
            os.fsync(f.fileno())
        os.replace(tmp, excel_path)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def _save_locked(excel_path, demo_name, framework, training, variant, dataset,
                 val_metrics, test_metrics, config, provenance=None):
    """Actual Excel write path — caller must hold the file lock."""
    import openpyxl

    def _get(d, *keys, default=None):
        """Try multiple keys, return first match."""
        for k in keys:
            try:
                v = d[k] if isinstance(d, dict) else getattr(d, k, None)
                if v is not None:
                    return v
            except (KeyError, TypeError):
                pass
        return default

    seed = _get(config, "seed")
    sheet_name = f"Seed_{seed}" if seed is not None else "Results"

    from openpyxl.styles import Font
    wb = _load_or_new_workbook(excel_path)

    # Get or create the target sheet
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
    else:
        ws = wb[sheet_name]
        # Extend the header of sheets created before provenance columns existed.
        header = [c.value for c in ws[1]]
        for col in COLUMNS:
            if col not in header:
                ws.cell(row=1, column=len(header) + 1, value=col).font = Font(bold=True)
                header.append(col)

    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        demo_name,
        framework,
        training,
        variant,
        dataset,
        # Validation
        round(float(val_metrics["loss"]), 6) if val_metrics else None,
        round(float(val_metrics["accuracy"]), 6) if val_metrics else None,
        round(float(val_metrics["auc"]), 6) if val_metrics else None,
        round(float(val_metrics["iou"]), 6) if val_metrics else None,
        round(float(val_metrics["dice"]), 6) if val_metrics else None,
        # Test
        round(float(test_metrics["loss"]), 6) if test_metrics else None,
        round(float(test_metrics["accuracy"]), 6) if test_metrics else None,
        round(float(test_metrics["auc"]), 6) if test_metrics else None,
        round(float(test_metrics["iou"]), 6) if test_metrics else None,
        round(float(test_metrics["dice"]), 6) if test_metrics else None,
        # Config
        _get(config, "partition", default="centralized"),
        _get(config, "dirichlet_alpha"),
        _get(config, "num_epochs", "num_train_steps", "fed_rounds", "num_train_steps"),
        _get(config, "num_clients"),
        _get(config, "local_epochs"),
        _get(config, "base_channels"),
        _get(config, "quantum_channels"),
        _get(config, "batch_size"),
        _get(config, "learning_rate"),
        None,  # Total_Params — optionally passed via config['total_params']
    ]

    # Allow callers to pass total_params via config
    if _get(config, "total_params") is not None:
        row[-1] = _get(config, "total_params")

    prov = provenance or _get(config, "provenance") or {}
    for _col, key in _PROVENANCE_COLUMN_KEYS:
        v = prov.get(key) if isinstance(prov, dict) else None
        row.append(None if v is None else str(v))

    # Place values by header name so older sheets with a different column
    # count still line up.
    header = [c.value for c in ws[1]]
    if header[:len(COLUMNS)] != COLUMNS:
        by_name = dict(zip(COLUMNS, row))
        row = [by_name.get(h) for h in header]

    ws.append(row)
    _atomic_save_workbook(wb, excel_path)
    print(f"Results saved to Excel: {excel_path}")
